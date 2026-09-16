"""Fold-local event inputs and separate intermediate/terminal supervision."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import torch
from sklearn.preprocessing import StandardScaler

from stageworld.artifacts import atomic_write_private_json, new_artifact_id, read_json
from stageworld.binary_data import parse_binary_label
from stageworld.data.baseline_clinical import (
    CT6_CLINICAL_SCHEMA,
    encode_baseline,
    fit_clinical_transform,
)
from stageworld.data.paired_ct import HMACPseudonymizer, _identifier
from stageworld.data.treatment_compact import encode_compact, fit_name_support
from stageworld.event_models import EventInputs
from stageworld.event_spec import ABSENT, CONFLICT, PRESENT, TASK, UNKNOWN
from stageworld.generated700_data import Pool, load_pool
from stageworld.synthetic_workflow import _atomic_torch_save


@dataclass
class EventPool:
    base: Pool
    events: torch.Tensor
    artifact_id: str

    def targets(self, rows: torch.Tensor, device: torch.device) -> dict[str, torch.Tensor]:
        valid = self.base.valid[rows].clone()
        valid[:, 0] &= self.events[rows, 1] == PRESENT
        return {
            "ct1": self.base.ct1_tokens[rows].to(device),
            "ct_valid": (self.base.ct0_valid & self.base.ct1_valid)[rows].to(device),
            "labels": self.base.labels[rows].to(device),
            "valid": valid.to(device),
        }

    def inputs(self, x: torch.Tensor, rows: torch.Tensor) -> EventInputs:
        return EventInputs(x[rows], self.base.ct0[rows], self.events[rows])


def event_status(surgery: int | None, postoperative: int | None) -> tuple[int, int]:
    operation = UNKNOWN if surgery is None else PRESENT if surgery else ABSENT
    if postoperative == 1:
        after = PRESENT if surgery == 1 else CONFLICT
    elif postoperative == 0:
        after = ABSENT
    else:
        after = UNKNOWN
    return operation, after


def prepare_pool(reference: Path, bindings_path: Path, output: Path) -> EventPool:
    import openpyxl  # type: ignore[import-untyped]
    from openpyxl.utils import column_index_from_string as ci  # type: ignore[import-untyped]

    base = load_pool(reference)
    if len(base.ids) != 651 or not base.ct0_valid.all():
        raise ValueError("Bind the authorized651 pool with available CT0")
    bindings = read_json(bindings_path)
    pseudo = HMACPseudonymizer.from_file(Path(bindings["STAGEWORLD_HMAC_KEY_FILE"]))
    workbook = openpyxl.load_workbook(
        bindings["STAGEWORLD_CLINICAL_EXCEL"], read_only=True, data_only=False
    )
    parsed: dict[str, tuple[int | None, int | None]] = {}
    labels: dict[str, list[int | None]] = {}
    audit: Counter[str] = Counter()
    try:
        sheet = workbook.worksheets[0]
        headers = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
        for column, expected in {
            "AQ": "\u662f\u5426\u884c\u80c3\u5207\u9664\u624b\u672f",
            "CE": "\u672f\u540e\u5316\u7597",
            "BJ": "pCR(1=\u662f\uff0c2=\u5426)",
            "CG": "\u590d\u53d1\u8f6c\u79fb\u72b6\u6001",
        }.items():
            if str(headers[ci(column) - 1]).strip() != expected:
                raise ValueError("Stage header binding differs")
        for row in sheet.iter_rows(min_row=3):
            key = _identifier(row[0].value, row[0].data_type)
            if key is None:
                audit["invalid_patient_key"] += 1
                continue
            patient = pseudo.token("patient", key, prefix="P")
            if patient in parsed:
                raise ValueError("Duplicate workbook patient")

            def binary(column: str, negative: int, row=row) -> int | None:
                cell = row[ci(column) - 1]
                return (
                    None
                    if cell.data_type in ("e", "f")
                    else parse_binary_label(cell.value, positive=1, negative=negative)
                )

            surgery = binary("AQ", 0)
            cell = row[ci("CE") - 1]
            text = str(cell.value).strip()
            postoperative = None
            if cell.data_type not in ("e", "f"):
                if text in ("\u662f", "\u6709", "1", "1.0"):
                    postoperative = 1
                elif text in ("\u5426", "\u65e0", "0", "0.0"):
                    postoperative = 0
            parsed[patient] = surgery, postoperative
            labels[patient] = [binary("BJ", 2), binary("CG", 0)]
            audit["workbook_patients"] += 1
            audit["pcr_missing"] += int(labels[patient][0] is None)
            audit["recurrence_missing"] += int(labels[patient][1] is None)
            audit["no_surgery_but_postoperative_yes"] += int(surgery == 0 and postoperative == 1)
    finally:
        workbook.close()
    events = []
    for index, patient in enumerate(base.ids):
        if patient not in parsed:
            raise ValueError("Stage data do not cover authorized patients")
        for endpoint in range(2):
            value = labels[patient][endpoint]
            valid = bool(base.valid[index, endpoint])
            if valid != (value is not None) or (valid and base.labels[index, endpoint] != value):
                raise ValueError("Current workbook labels differ from frozen pool")
        methods = list(base.treatments[patient]["methods"].values())
        neoadjuvant = (
            PRESENT if 1 in methods else ABSENT if all(v == 0 for v in methods) else UNKNOWN
        )
        events.append([neoadjuvant, *event_status(*parsed[patient])])
    tensor = torch.tensor(events, dtype=torch.long)
    output.mkdir(parents=True, exist_ok=True, mode=0o700)
    path = output / "events.pt"
    if path.exists():
        saved = torch.load(path, weights_only=True, map_location="cpu")
        if saved["source_pool_id"] != base.artifact_id or not torch.equal(saved["events"], tensor):
            raise ValueError("Event source changed during an existing run")
        artifact_id = saved["artifact_id"]
    else:
        artifact_id = new_artifact_id("event-trajectory-inputs")
        _atomic_torch_save(
            path,
            {
                "artifact_id": artifact_id,
                "source_pool_id": base.artifact_id,
                "patient_ids": base.ids,
                "events": tensor,
                "task": TASK,
            },
        )
    atomic_write_private_json(
        output / "audit.json",
        {
            "source_workbook": dict(audit),
            "selected_patients": len(base.ids),
            "event_status_counts": [
                torch.bincount(tensor[:, i], minlength=4).tolist() for i in range(3)
            ],
            "pcr": {
                "valid": int(base.valid[:, 0].sum()),
                "positive": int(base.labels[:, 0][base.valid[:, 0]].sum()),
            },
            "recurrence": {
                "valid": int(base.valid[:, 1].sum()),
                "positive": int(base.labels[:, 1][base.valid[:, 1]].sum()),
            },
            "time_and_cycles_used": False,
            "postoperative_policy": "recorded_yes_accepted_only_after_confirmed_surgery",
        },
    )
    return EventPool(base, tensor, artifact_id)


def encode_inputs(clinical: list[dict], treatments: list[dict], snapshot: dict) -> torch.Tensor:
    baseline = encode_baseline(clinical, snapshot["clinical"])
    actions, _ = encode_compact(treatments, snapshot["support"])
    raw = torch.cat((baseline.ridge_features(), actions.flatten(1)), 1).double()
    if "mean" not in snapshot:
        return raw
    result = ((raw - snapshot["mean"]) / snapshot["scale"]).float()
    if not torch.isfinite(result).all():
        raise ValueError("Nonfinite event-only tabular inputs")
    return result


def fit_inputs(pool: EventPool, fitting: list[str], path: Path) -> tuple[torch.Tensor, dict]:
    base = pool.base
    rows = base.indices(fitting)
    if not fitting or len(set(fitting)) != len(fitting):
        raise ValueError("Require unique training patients")
    clinical = [base.clinical[p] for p in base.ids]
    treatment = [base.treatments[p] for p in base.ids]
    if path.exists():
        snapshot = torch.load(path, weights_only=True, map_location="cpu")
        if snapshot["fit_ids"] != fitting or snapshot["pool_id"] != pool.artifact_id:
            raise ValueError("Training transform binding differs")
    else:
        snapshot = {
            "artifact_id": new_artifact_id("event-input-transform"),
            "task": TASK,
            "pool_id": pool.artifact_id,
            "fit_ids": fitting,
            "clinical": fit_clinical_transform(
                base.clinical, set(fitting), schema_version=CT6_CLINICAL_SCHEMA
            ),
            "support": fit_name_support(treatment, set(fitting)),
        }
        raw = encode_inputs(clinical, treatment, snapshot)
        scaler = StandardScaler().fit(raw[rows].numpy())
        snapshot.update(mean=torch.from_numpy(scaler.mean_), scale=torch.from_numpy(scaler.scale_))
        if not np.isfinite(scaler.scale_).all():
            raise ValueError("Invalid fitted input scale")
        _atomic_torch_save(path, snapshot)
    return encode_inputs(clinical, treatment, snapshot), snapshot
