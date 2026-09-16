"""Private, outcome-independent treatment-modality summaries for paired CT transitions."""

from __future__ import annotations

import math
import unicodedata
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import replace
from pathlib import Path
from typing import Any

import torch

from stageworld.artifacts import atomic_write_private_json, new_artifact_id, read_json
from stageworld.config import StageWorldConfig
from stageworld.data.paired_ct import HMACPseudonymizer, _identifier
from stageworld.data.treatment_regimens import (
    REGIMEN_ACTION_DIM,
    REGIMEN_PROTOCOLS,
    REGIMEN_SCHEMA,
    encode_regimens,
    fit_cycle_transform,
    read_regimen_rows,
    summarize_regimens,
)
from stageworld.errors import ArtifactError, DataContractError
from stageworld.model import ActionTokens
from stageworld.real_workflow import RealFeatureBundle
from stageworld.training import WorldModelBatch

TREATMENT_PROTOCOL = "gastric-roi-treatment-os-development-v1"
TREATMENT_SCHEMA = "paired-ct-treatment-modality-summary-v1"
CONFIRMED_CUTOFF = "completed_before_post_treatment_ct"
# Codes differ by column; a zero is NOT the documented negative code for every field.
TREATMENT_FIELDS = (
    ("chemotherapy", "AA", 0, "\u5316\u7597(1=\u6709,0=\u65e0)"),
    ("immunotherapy", "AB", 2, "\u514d\u75ab(1=\u6709,2=\u65e0)"),
    ("targeted", "AC", 2, "\u9776\u5411(1=\u6709,2=\u65e0)"),
    ("interventional", "AG", 0, "\u4ecb\u5165(1=\u6709,0=\u65e0)"),
    ("hipec", "AH", 2, "HIPEC(1=\u6709,2=\u65e0)"),
)
METHOD_NAMES = tuple(row[0] for row in TREATMENT_FIELDS)


def parse_treatment_flag(value: Any, data_type: str, negative_code: int) -> int | None:
    if data_type in {"e", "f"} or value is None or isinstance(value, bool):
        return None
    if isinstance(value, str) and value.strip() not in {"0", "1", "2"}:
        return None
    if not isinstance(value, (str, int, float)):
        return None
    number = float(value)
    if not math.isfinite(number):
        return None
    if number == 1:
        return 1
    return 0 if number == negative_code else None


def _header(value: Any) -> str:
    return "".join(unicodedata.normalize("NFKC", str(value)).split())


def read_treatment_rows(
    workbook_path: Path,
    *,
    patient_ids: set[str],
    pseudonymizer: HMACPseudonymizer,
) -> list[dict[str, Any]]:
    import openpyxl  # type: ignore[import-untyped]
    from openpyxl.utils import column_index_from_string  # type: ignore[import-untyped]

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        headers = next(sheet.iter_rows(min_row=1, max_row=1, max_col=34))
        indices = {name: column_index_from_string(col) - 1 for name, col, _, _ in TREATMENT_FIELDS}
        if any(
            _header(headers[indices[name]].value) != _header(expected)
            for name, _, _, expected in TREATMENT_FIELDS
        ):
            raise DataContractError(
                code="TREATMENT_HEADER_MISMATCH", message="Treatment column coding changed."
            )
        result: dict[str, dict[str, Any]] = {}
        # No surgery, pathology, survival or follow-up columns enter this reader.
        for cells in sheet.iter_rows(min_row=2, max_col=34):
            key = _identifier(cells[0].value, cells[0].data_type)
            if key is None:
                continue
            patient = pseudonymizer.token("patient", key, prefix="P")
            if patient not in patient_ids:
                continue
            if patient in result:
                raise DataContractError(
                    code="TREATMENT_DUPLICATE_PATIENT", message="Duplicate treatment row."
                )
            methods = {}
            for name, _, negative_code, _ in TREATMENT_FIELDS:
                cell = cells[indices[name]]
                methods[name] = parse_treatment_flag(cell.value, cell.data_type, negative_code)
            result[patient] = {"patient_id": patient, "methods": methods}
        if set(result) != patient_ids:
            raise DataContractError(
                code="TREATMENT_PATIENT_MISSING", message="A development treatment row is missing."
            )
        return [result[p] for p in sorted(result)]
    finally:
        workbook.close()


def summarize_treatments(rows: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    return {
        "patients": len(rows),
        "modalities": {
            name: dict(
                Counter(
                    "unknown"
                    if r["methods"][name] is None
                    else "present"
                    if r["methods"][name] == 1
                    else "absent"
                    for r in rows
                )
            )
            for name in METHOD_NAMES
        },
        "drug_regimens_used": False,
        "cycle_counts_used": False,
        "surgery_or_pathology_used": False,
        "outcomes_used": False,
        "test_used": False,
    }


def require_treatment_cutoff(config: StageWorldConfig) -> None:
    if config.clinical.treatment_summary_cutoff != CONFIRMED_CUTOFF:
        raise DataContractError(
            code="TREATMENT_CT_CUTOFF_UNCONFIRMED",
            message="Confirm that included treatments were completed before the target CT.",
        )


def prepare_treatments(
    config: StageWorldConfig,
    bundle: RealFeatureBundle,
    pseudonymizer: HMACPseudonymizer,
    *,
    audit_only: bool = False,
) -> dict[str, Any]:
    if not config.paths.clinical_excel:
        raise DataContractError(code="TREATMENT_WORKBOOK_MISSING", message="Bind the workbook.")
    patients = {
        p for batches in bundle.batches_by_split.values() for b in batches for p in b.patient_ids
    }
    regimen_mode = config.training.development_protocol in REGIMEN_PROTOCOLS
    reader = read_regimen_rows if regimen_mode else read_treatment_rows
    rows = reader(
        Path(config.paths.clinical_excel), patient_ids=patients, pseudonymizer=pseudonymizer
    )
    aggregate = summarize_regimens(rows) if regimen_mode else summarize_treatments(rows)
    if audit_only:
        return {
            "status": "audited_not_trained",
            **aggregate,
            "cutoff_confirmed": config.clinical.treatment_summary_cutoff == CONFIRMED_CUTOFF,
        }
    require_treatment_cutoff(config)
    if not config.paths.treatment_manifest:
        raise DataContractError(
            code="TREATMENT_MANIFEST_MISSING", message="Bind a private manifest."
        )
    path = Path(config.paths.treatment_manifest)
    payload = {
        "schema_version": REGIMEN_SCHEMA if regimen_mode else TREATMENT_SCHEMA,
        "artifact_id": new_artifact_id("treatment-interval"),
        "data_lineage_id": bundle.data_lineage_id,
        "cohort_artifact_id": bundle.cohort_artifact_id,
        "split_version": bundle.split_version,
        "ct_feature_artifact_id": bundle.feature_artifact_id,
        "cutoff": CONFIRMED_CUTOFF,
        "time_basis": "confirmed_interval_summary_at_target_ct_not_a_dosing_timestamp",
        "rows": rows,
        "aggregate": aggregate,
    }
    if regimen_mode:
        training_ids = {p for b in bundle.batches_by_split["train"] for p in b.patient_ids}
        payload["cycle_transform"] = fit_cycle_transform(rows, training_ids)
        payload["cycle_transform_fit_split"] = "train"
        payload["cutoff_confirmation"] = "user_2026-09-10_stage2_through_post_ct"
    if path.exists():
        existing = read_json(path)
        if {k: v for k, v in existing.items() if k != "artifact_id"} != {
            k: v for k, v in payload.items() if k != "artifact_id"
        }:
            raise ArtifactError(
                code="TREATMENT_MANIFEST_CHANGED", message="Use a new treatment manifest version."
            )
        payload = existing
    else:
        version = path.parent / "versions" / f"{payload['artifact_id']}.json"
        atomic_write_private_json(version, payload)
        atomic_write_private_json(path, payload)
    return {"status": "prepared", "treatment_artifact_id": payload["artifact_id"], **aggregate}


def encode_treatments(
    batch: WorldModelBatch,
    rows: Mapping[str, Mapping[str, Any]],
) -> WorldModelBatch:
    values = torch.zeros(batch.batch_size, 5, 8)
    for i, patient in enumerate(batch.patient_ids):
        methods = rows[patient]["methods"]
        if set(methods) != set(METHOD_NAMES):
            raise ArtifactError(
                code="TREATMENT_FIELDS_CHANGED", message="Unexpected method fields."
            )
        for j, name in enumerate(METHOD_NAMES):
            value = methods[name]
            if value is not None and (type(value) is not int or value not in (0, 1)):
                raise ArtifactError(
                    code="TREATMENT_VALUE_INVALID", message="Invalid treatment flag."
                )
            values[i, j, j] = 1
            values[i, j, 5] = 0 if value is None else value
            values[i, j, 6] = float(value is not None)
            values[i, j, 7] = 1  # Interval-summary marker, not a per-cycle dose.
    times = batch.ct1_acquisition_time[:, None].expand(-1, 5).clone()
    actions = ActionTokens(
        values=values,
        valid=torch.ones(batch.batch_size, 5, dtype=torch.bool),
        event_time=times,
        available_time=times.clone(),
        event_type=torch.full((batch.batch_size, 5), 5, dtype=torch.long),
        planned_or_delivered=torch.ones(batch.batch_size, 5, dtype=torch.long),
        provenance=(TREATMENT_SCHEMA, "confirmed_pre_ct_interval_summary", "not_dosing_times"),
    )
    actions.validate()
    return replace(batch, treatment_actions=actions)


def attach_treatments(config: StageWorldConfig, bundle: RealFeatureBundle) -> RealFeatureBundle:
    require_treatment_cutoff(config)
    regimen_mode = config.training.development_protocol in REGIMEN_PROTOCOLS
    expected_dim = REGIMEN_ACTION_DIM if regimen_mode else 8
    if config.model.action_input_dim != expected_dim or not config.paths.treatment_manifest:
        raise ArtifactError(code="TREATMENT_CONFIG_INVALID", message="Check treatment bindings.")
    path = Path(config.paths.treatment_manifest)
    if path.stat().st_mode & 0o077:
        raise ArtifactError(code="TREATMENT_PERMISSIONS_UNSAFE", message="Manifest is not private.")
    payload = read_json(path)
    expected = {
        "schema_version": REGIMEN_SCHEMA if regimen_mode else TREATMENT_SCHEMA,
        "cutoff": CONFIRMED_CUTOFF,
        "data_lineage_id": bundle.data_lineage_id,
        "cohort_artifact_id": bundle.cohort_artifact_id,
        "split_version": bundle.split_version,
        "ct_feature_artifact_id": bundle.feature_artifact_id,
    }
    if any(payload.get(k) != v for k, v in expected.items()):
        raise ArtifactError(code="TREATMENT_LINEAGE_MISMATCH", message="Treatment lineage changed.")
    version = path.parent / "versions" / f"{payload['artifact_id']}.json"
    if read_json(version) != payload:
        raise ArtifactError(
            code="TREATMENT_SNAPSHOT_MISMATCH", message="Treatment snapshot changed."
        )
    rows = {r["patient_id"]: r for r in payload["rows"]}
    patients = {
        p for batches in bundle.batches_by_split.values() for b in batches for p in b.patient_ids
    }
    if set(rows) != patients or len(rows) != len(payload["rows"]):
        raise ArtifactError(
            code="TREATMENT_POPULATION_MISMATCH", message="Treatment patients differ."
        )
    if regimen_mode:
        training_ids = {p for b in bundle.batches_by_split["train"] for p in b.patient_ids}
        if payload.get("cycle_transform_fit_split") != "train" or payload.get(
            "cycle_transform"
        ) != fit_cycle_transform(payload["rows"], training_ids):
            raise ArtifactError(
                code="CYCLE_TRANSFORM_MISMATCH", message="Cycle transform is not training-only."
            )
    return replace(
        bundle,
        treatment_snapshot=payload,
        batches_by_split={
            split: tuple(
                encode_regimens(b, rows, payload["cycle_transform"])
                if regimen_mode
                else encode_treatments(b, rows)
                for b in batches
            )
            for split, batches in bundle.batches_by_split.items()
        },
    )
