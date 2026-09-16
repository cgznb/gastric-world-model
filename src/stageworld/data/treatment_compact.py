"""Four interval-condition tokens retaining named drugs, without cycle inputs."""

from __future__ import annotations

from collections import Counter
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

import torch
from torch import Tensor

from stageworld.data.paired_ct import HMACPseudonymizer, _identifier
from stageworld.data.treatment_regimens import (
    DRUG_NAMES,
    GROUPED_COLUMNS,
    MENTION_NAMES,
    REGIMEN_CODES,
    parse_named_mentions,
)
from stageworld.data.treatment_summary import (
    METHOD_NAMES,
    TREATMENT_FIELDS,
    _header,
    parse_treatment_flag,
)
from stageworld.errors import DataContractError
from stageworld.model.types import ActionTokens

COMPACT_TREATMENT_SCHEMA = "paired-ct-regimen-drug-no-cycle-v1"
DRUG_CODES = tuple(name for name, _ in DRUG_NAMES)
DESCRIPTOR_NAMES = (*METHOD_NAMES, *MENTION_NAMES)
COMPACT_GROUPS = (
    tuple(range(5)),
    tuple(range(5, 23)),
    tuple(range(23, 33)),
    tuple(range(33, 39)),
)
COMPACT_ACTION_DIM = 82
FIELD_KEYS = {"methods", "regimens", "drugs"}


def normalize_treatment(
    row: Mapping[str, Any],
    *,
    strict_conflicts: bool = True,
) -> dict[str, Any]:
    if not isinstance(row, Mapping) or set(row) != FIELD_KEYS:
        raise DataContractError(code="COMPACT_TREATMENT_FIELDS", message="Use no-cycle fields.")
    result: dict[str, Any] = {}
    for key, names in (
        ("methods", METHOD_NAMES),
        ("regimens", REGIMEN_CODES),
        ("drugs", DRUG_CODES),
    ):
        values = row[key]
        if not isinstance(values, Mapping) or not set(values).issubset(names):
            raise DataContractError(
                code="COMPACT_TREATMENT_NAME", message="Unknown treatment name."
            )
        if any(v is not None and (type(v) is not int or v not in (0, 1)) for v in values.values()):
            raise DataContractError(code="COMPACT_TREATMENT_VALUE", message="Use 0, 1 or null.")
        result[key] = {name: values.get(name) for name in names}
    conflicts = []
    for method, drugs in (
        ("chemotherapy", DRUG_CODES[:9]),
        ("immunotherapy", DRUG_CODES[9:19]),
        ("targeted", DRUG_CODES[19:]),
    ):
        named = any(result["drugs"][n] == 1 for n in drugs)
        if method == "chemotherapy":
            named |= any(v == 1 for v in result["regimens"].values())
        if result["methods"][method] == 0 and named:
            conflicts.append(method)
            result["methods"][method] = None
            for name in drugs:
                result["drugs"][name] = None
            if method == "chemotherapy":
                result["regimens"] = dict.fromkeys(REGIMEN_CODES)
    if conflicts and strict_conflicts:
        raise DataContractError(
            code="COMPACT_TREATMENT_CONFLICT", message="Conflicting conditions."
        )
    return {**result, "conflicts": conflicts}


def treatment_fields(row: Mapping[str, Any]) -> dict[str, Any]:
    return {key: row[key] for key in ("methods", "regimens", "drugs")}


def descriptors(row: Mapping[str, Any]) -> list[int | None]:
    clean = normalize_treatment(treatment_fields(row))
    return (
        [clean["methods"][n] for n in METHOD_NAMES]
        + [clean["regimens"][n] for n in REGIMEN_CODES]
        + [clean["drugs"][n] for n in DRUG_CODES]
    )


def read_compact_rows(
    path: Path,
    patient_ids: set[str],
    pseudonymizer: HMACPseudonymizer,
) -> list[dict[str, Any]]:
    import openpyxl  # type: ignore[import-untyped]
    from openpyxl.utils import column_index_from_string as ci  # type: ignore[import-untyped]

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        groups = next(sheet.iter_rows(min_row=1, max_row=1, max_col=38))
        headers = next(sheet.iter_rows(min_row=2, max_row=2, max_col=38))
        expected = {"A": "\u5e8f\u5217\u53f7", "AD": "\u672f\u524d\u6cbb\u7597"}
        expected.update({c: f[3] for c, f in zip(GROUPED_COLUMNS, TREATMENT_FIELDS, strict=True)})
        if _header(groups[ci("AD") - 1].value) != _header(
            "\u9636\u6bb52\u2014\u2014\u672f\u524d\u5316\u7597\u65b9\u6848\u3001\u5468\u671f\u53caCT\u8bc4\u4f30"
        ) or any(_header(headers[ci(c) - 1].value) != _header(h) for c, h in expected.items()):
            raise DataContractError(code="COMPACT_HEADER_MISMATCH", message="Workbook differs.")
        result = {}
        for cells in sheet.iter_rows(min_row=3, max_col=38):
            key = _identifier(cells[0].value, cells[0].data_type)
            if key is None:
                continue
            patient = pseudonymizer.token("patient", key, prefix="P")
            if patient not in patient_ids:
                continue
            if patient in result:
                raise DataContractError(
                    code="TREATMENT_DUPLICATE_PATIENT", message="Duplicate row."
                )
            methods = {}
            for col, (name, _, negative, _) in zip(GROUPED_COLUMNS, TREATMENT_FIELDS, strict=True):
                cell = cells[ci(col) - 1]
                methods[name] = parse_treatment_flag(cell.value, cell.data_type, negative)
            summary = cells[ci("AD") - 1]
            mentions = parse_named_mentions(summary.value, summary.data_type)
            row = normalize_treatment(
                {
                    "methods": methods,
                    "regimens": {c: mentions[f"regimen_{c.lower()}"] for c in REGIMEN_CODES},
                    "drugs": {n: mentions[n] for n in DRUG_CODES},
                },
                strict_conflicts=False,
            )
            result[patient] = {"patient_id": patient, **row}
        if set(result) != patient_ids:
            raise DataContractError(code="TREATMENT_PATIENT_MISSING", message="Incomplete rows.")
        return [result[p] for p in sorted(result)]
    finally:
        workbook.close()


def fit_name_support(rows: Sequence[Mapping[str, Any]], training_ids: set[str]) -> dict[str, Any]:
    if not training_ids or not training_ids <= {r["patient_id"] for r in rows}:
        raise DataContractError(code="TREATMENT_SUPPORT_SPLIT", message="Bind training patients.")
    counts = dict.fromkeys(MENTION_NAMES, 0)
    for row in rows:
        if row["patient_id"] in training_ids:
            for name, value in zip(MENTION_NAMES, descriptors(row)[5:], strict=True):
                counts[name] += int(value == 1)
    return {
        "schema_version": COMPACT_TREATMENT_SCHEMA,
        "fit_split": "train",
        "training_patients": len(training_ids),
        "positive_counts": counts,
        "descriptor_names": list(DESCRIPTOR_NAMES),
        "descriptor_groups": [list(group) for group in COMPACT_GROUPS],
    }


def encode_compact(
    rows: Sequence[Mapping[str, Any]],
    support: Mapping[str, Any],
) -> tuple[Tensor, tuple[tuple[str, ...], ...]]:
    counts = support.get("positive_counts", {})
    if (
        support.get("schema_version") != COMPACT_TREATMENT_SCHEMA
        or support.get("fit_split") != "train"
        or support.get("descriptor_names") != list(DESCRIPTOR_NAMES)
        or support.get("descriptor_groups") != [list(group) for group in COMPACT_GROUPS]
        or set(counts) != set(MENTION_NAMES)
        or any(type(n) is not int or n < 0 for n in counts.values())
        or not rows
    ):
        raise DataContractError(code="TREATMENT_SUPPORT_INVALID", message="Bind training support.")
    values = torch.zeros(len(rows), 4, COMPACT_ACTION_DIM)
    flags = []
    for i, row in enumerate(rows):
        fields = descriptors(row)
        unseen = []
        for j, name in enumerate(MENTION_NAMES, start=5):
            if counts[name] == 0:
                if fields[j] == 1:
                    unseen.append(f"unseen_treatment_name:{name}")
                fields[j] = None
        for g, indices in enumerate(COMPACT_GROUPS):
            values[i, g, 78 + g] = 1
            for j in indices:
                value = fields[j]
                values[i, g, j] = 0 if value is None else value
                values[i, g, 39 + j] = float(value is not None)
        flags.append(tuple(unseen))
    return values, tuple(flags)


def compact_actions(values: Tensor, target_time: Tensor, *, provenance: str) -> ActionTokens:
    if values.shape != (target_time.shape[0], 4, COMPACT_ACTION_DIM):
        raise DataContractError(code="COMPACT_ACTION_SHAPE", message="Use four treatment tokens.")
    times = target_time[:, None].expand(-1, 4).clone()
    actions = ActionTokens(
        values=values,
        valid=torch.ones_like(times, dtype=torch.bool),
        event_time=times,
        available_time=times.clone(),
        event_type=torch.full_like(times, 5, dtype=torch.long),
        planned_or_delivered=torch.ones_like(times, dtype=torch.long),
        provenance=(COMPACT_TREATMENT_SCHEMA, provenance, "not_dosing_times"),
    )
    actions.validate()
    return actions


def summarize_compact(
    rows: Sequence[Mapping[str, Any]], support: Mapping[str, Any]
) -> dict[str, Any]:
    _, flags = encode_compact(rows, support)
    return {
        "schema_version": COMPACT_TREATMENT_SCHEMA,
        "patients": len(rows),
        "descriptor_count": 39,
        "action_token_count": 4,
        "action_input_dim": COMPACT_ACTION_DIM,
        "named_drug_count": len(DRUG_CODES),
        "cycle_counts_used": False,
        "raw_text_used_as_model_input": False,
        "outcomes_used": False,
        "test_used": False,
        "positive_counts": {
            n: sum(descriptors(r)[i] == 1 for r in rows) for i, n in enumerate(DESCRIPTOR_NAMES)
        },
        "unknown_counts": {
            n: sum(descriptors(r)[i] is None for r in rows) for i, n in enumerate(DESCRIPTOR_NAMES)
        },
        "conflicts": dict(Counter(c for r in rows for c in r.get("conflicts", []))),
        "unseen_name_flags": dict(Counter(f for patient in flags for f in patient)),
        "support_fit_split": "train",
        "missing_name_means": "not_documented_not_absence",
        "time_basis": "retrospective_interval_summary_not_baseline_known_fact",
    }
