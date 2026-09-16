"""Static, explicitly baseline clinical fields for generated-state prognosis."""

from __future__ import annotations

import math
import re
import statistics
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import torch
from torch import Tensor

from stageworld.data.audit import ExcelAuditConfig, _semantic_missing
from stageworld.data.paired_ct import HMACPseudonymizer, _identifier
from stageworld.data.treatment_summary import _header
from stageworld.errors import DataContractError

BASELINE_CLINICAL_SCHEMA = "gastric-baseline-19-fields-v2"


@dataclass(frozen=True)
class ClinicalField:
    name: str
    column: str
    header: str
    categories: tuple[str, ...] = ()


CLINICAL_FIELDS = (
    ClinicalField("sex", "F", "性别", ("male", "female")),
    ClinicalField("age", "H", "年龄"),
    ClinicalField("height", "I", "身高（cm）"),
    ClinicalField("weight", "J", "体重"),
    ClinicalField("bmi", "K", "BMI"),
    ClinicalField("lauren", "L", "Lauren分型-基线", ("1", "2", "3", "4")),
    ClinicalField("signet", "M", "印戒成分-基线", ("0", "1")),
    ClinicalField("differentiation", "N", "分化程度-基线", tuple(map(str, range(1, 7)))),
    ClinicalField("her2", "O", "HER2-基线", ("0", "1")),
    ClinicalField("mmr", "P", "MMR-基线", ("0", "1", "2", "3")),
    ClinicalField("pdl1", "Q", "PD-L1-基线", ("0", "1")),
    ClinicalField("tps", "R", "TPS(%)-基线"),
    ClinicalField("cps", "S", "CPS-基线"),
    ClinicalField("eber", "T", "EBER-基线", ("0", "1")),
    ClinicalField("location", "U", "肿瘤位置", ("1", "2")),
    ClinicalField("dentate_line", "V", "是否累及齿状线", ("1", "2")),
    ClinicalField(
        "ct_stage",
        "W",
        "基线cT分期",
        (
            "0",
            "is",
            "1",
            "1a",
            "1b",
            "2",
            "3",
            "4",
            "4a",
            "4b",
            "x",
        ),
    ),
    ClinicalField("cn_stage", "X", "基线cN分期", ("0", "1", "2", "3", "3a", "3b", "x", "+")),
    ClinicalField("cm_stage", "Y", "基线cM分期", ("0", "1", "x")),
)
FIELD_NAMES = tuple(field.name for field in CLINICAL_FIELDS)
CT6_CLINICAL_SCHEMA = "gastric-baseline-6-fields-v1"
CT6_FIELD_NAMES = ("sex", "age", "bmi", "ct_stage", "cn_stage", "cm_stage")
CT6_FIELDS = tuple(next(f for f in CLINICAL_FIELDS if f.name == n) for n in CT6_FIELD_NAMES)
ClinicalValue = float | str | None


def clinical_fields(schema_version: str) -> tuple[ClinicalField, ...]:
    if schema_version == BASELINE_CLINICAL_SCHEMA:
        return CLINICAL_FIELDS
    if schema_version == CT6_CLINICAL_SCHEMA:
        return CT6_FIELDS
    raise DataContractError(code="CLINICAL_SCHEMA_UNKNOWN", message="Unknown clinical schema.")


def parse_clinical_value(field: ClinicalField, value: Any) -> tuple[ClinicalValue, str]:
    if _semantic_missing(value, "", ExcelAuditConfig().missing_tokens):
        return None, "missing"
    text = str(value).strip().casefold()
    if text in {"未测", "未检测", "未做", "不详", "未知", "-", "/"}:
        return None, "not_recorded"
    if field.name == "sex":
        normalized = {"男": "male", "女": "female", "m": "male", "f": "female"}
        text = normalized.get(text, text)
    if field.name in {"ct_stage", "cn_stage", "cm_stage"}:
        letter = field.name[1]
        text = re.sub(rf"^c?{letter}", "", text)
    if field.name == "tps" and text.endswith("%"):
        text = text[:-1].strip()
    try:
        numeric = float(text)
    except (ValueError, TypeError):
        numeric = None
    if numeric is not None and not math.isfinite(numeric):
        return None, "nonfinite"
    if field.categories:
        if numeric is not None and numeric.is_integer():
            text = str(int(numeric))
        if field.name == "lauren" and text == "5":
            return None, "not_measured"
        return (text, "observed") if text in field.categories else (None, "invalid_category")
    if numeric is None:
        return None, "invalid_numeric"
    if numeric < 0 or (field.name == "tps" and numeric > 100):
        return None, "outside_field_range"
    return numeric, "observed"


def parse_baseline_fields(
    values: Mapping[str, Any],
    *,
    schema_version: str = BASELINE_CLINICAL_SCHEMA,
) -> dict[str, ClinicalValue]:
    selected = clinical_fields(schema_version)
    if not isinstance(values, Mapping) or not set(values).issubset(f.name for f in selected):
        raise DataContractError(
            code="BASELINE_FIELD_NOT_ALLOWED", message="Use only the declared baseline fields."
        )
    return {
        field.name: parse_clinical_value(field, values.get(field.name))[0] for field in selected
    }


def read_baseline_rows(
    workbook_path: Path,
    patient_ids: set[str],
    pseudonymizer: HMACPseudonymizer,
    *,
    schema_version: str = BASELINE_CLINICAL_SCHEMA,
) -> tuple[dict[str, dict[str, ClinicalValue]], dict[str, Any]]:
    import openpyxl  # type: ignore[import-untyped]
    from openpyxl.utils import column_index_from_string as ci  # type: ignore[import-untyped]

    selected = clinical_fields(schema_version)
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        groups = next(sheet.iter_rows(min_row=1, max_row=1, max_col=25, values_only=True))
        headers = next(sheet.iter_rows(min_row=2, max_row=2, max_col=25, values_only=True))
        if (
            _header(groups[4]) != _header("阶段1——患者基线数据")
            or _header(headers[0]) != _header("序列号")
            or any(
                not _header(headers[ci(f.column) - 1]).startswith(_header(f.header))
                for f in selected
            )
        ):
            raise DataContractError(
                code="BASELINE_HEADER_MISMATCH", message="Expected the grouped baseline workbook."
            )
        rows: dict[str, dict[str, ClinicalValue]] = {}
        quality = {field.name: Counter[str]() for field in selected}
        # The row iterator stops at Y: ambiguous Z and every later-stage value are unread.
        for cells in sheet.iter_rows(min_row=3, max_col=25):
            key = _identifier(cells[0].value, cells[0].data_type)
            if key is None:
                continue
            patient = pseudonymizer.token("patient", key, prefix="P")
            if patient not in patient_ids:
                continue
            if patient in rows:
                raise DataContractError(
                    code="DUPLICATE_BASELINE_PATIENT", message="Repeated baseline patient."
                )
            row: dict[str, ClinicalValue] = {}
            for field in selected:
                cell = cells[ci(field.column) - 1]
                value, reason = (
                    (None, "error_or_formula")
                    if cell.data_type in {"e", "f"}
                    else parse_clinical_value(field, cell.value)
                )
                row[field.name] = value
                quality[field.name][reason] += 1
            rows[patient] = row
        if set(rows) != patient_ids:
            raise DataContractError(
                code="BASELINE_PATIENT_MISSING", message="Baseline rows must cover retained pairs."
            )
        return rows, {
            "schema_version": schema_version,
            "patients": len(rows),
            "fields": {name: dict(counts) for name, counts in quality.items()},
            "source_columns": {f.name: f.column for f in selected},
            "availability": "declared_pre_treatment_stage_no_per_field_timestamps",
            "ambiguous_lauren_Z_used": False,
            "outcomes_used": False,
            "test_used": False,
        }
    finally:
        workbook.close()


def fit_clinical_transform(
    rows: Mapping[str, Mapping[str, ClinicalValue]],
    training_ids: set[str],
    *,
    schema_version: str = BASELINE_CLINICAL_SCHEMA,
) -> dict[str, Any]:
    if not training_ids or not training_ids <= rows.keys():
        raise DataContractError(code="CLINICAL_FIT_SPLIT", message="Bind training patients only.")
    fields: dict[str, Any] = {}
    for field in clinical_fields(schema_version):
        if field.categories:
            continue
        values = [
            float(value) for p in sorted(training_ids) if (value := rows[p][field.name]) is not None
        ]
        mean = statistics.fmean(values) if values else 0.0
        scale = statistics.pstdev(values) if len(values) > 1 else 0.0
        fields[field.name] = {
            "mean": mean,
            "scale": scale if scale > 1e-6 else 1.0,
            "n": len(values),
        }
    return {
        "schema_version": schema_version,
        "fit_split": "train",
        "training_patients": len(training_ids),
        "continuous": fields,
    }


@dataclass(frozen=True)
class BaselineClinical:
    values: Tensor
    categories: Tensor
    observed: Tensor
    schema_version: str = BASELINE_CLINICAL_SCHEMA

    def validate(self) -> None:
        selected = clinical_fields(self.schema_version)
        shape = self.values.shape
        if (
            len(shape) != 2
            or shape[1] != len(selected)
            or self.categories.shape != shape
            or self.observed.shape != shape
            or self.observed.dtype is not torch.bool
            or self.categories.dtype is not torch.long
            or self.categories.device != self.values.device
            or self.observed.device != self.values.device
            or not torch.isfinite(self.values).all()
            or self.values.requires_grad
        ):
            raise DataContractError(
                code="CLINICAL_TENSOR_CONTRACT", message="Invalid baseline data."
            )
        for i, field in enumerate(selected):
            codes = self.categories[:, i]
            if (codes < 0).any() or (codes > len(field.categories)).any():
                raise DataContractError(code="CLINICAL_CATEGORY_RANGE", message="Invalid category.")
            if field.categories and not torch.equal(codes > 0, self.observed[:, i]):
                raise DataContractError(
                    code="CLINICAL_CATEGORY_MASK", message="Category mask differs."
                )
        if (self.values[~self.observed] != 0).any():
            raise DataContractError(
                code="CLINICAL_MISSING_VALUE", message="Missing values use zero."
            )

    def to(self, device: torch.device | str) -> BaselineClinical:
        return BaselineClinical(
            self.values.to(device),
            self.categories.to(device),
            self.observed.to(device),
            self.schema_version,
        )

    def ridge_features(self) -> Tensor:
        self.validate()
        columns = []
        for i, field in enumerate(clinical_fields(self.schema_version)):
            if field.categories:
                columns.append(
                    torch.nn.functional.one_hot(
                        self.categories[:, i],
                        len(field.categories) + 1,
                    ).float()
                )
            else:
                columns.extend((self.values[:, i : i + 1], self.observed[:, i : i + 1].float()))
        return torch.cat(columns, dim=1)


def encode_baseline(
    rows: Sequence[Mapping[str, ClinicalValue]],
    transform: Mapping[str, Any],
) -> BaselineClinical:
    schema_version = str(transform.get("schema_version", ""))
    selected = clinical_fields(schema_version)
    if not rows or transform.get("fit_split") != "train":
        raise DataContractError(
            code="CLINICAL_TRANSFORM_INVALID", message="Bind fitted baseline data."
        )
    values = torch.zeros(len(rows), len(selected))
    categories = torch.zeros_like(values, dtype=torch.long)
    observed = torch.zeros_like(values, dtype=torch.bool)
    for row_index, row in enumerate(rows):
        if set(row) != {f.name for f in selected}:
            raise DataContractError(code="CLINICAL_FIELD_SCHEMA", message="Baseline fields differ.")
        for index, field in enumerate(selected):
            value = row[field.name]
            if value is None:
                continue
            observed[row_index, index] = True
            if field.categories:
                if str(value) not in field.categories:
                    raise DataContractError(
                        code="CLINICAL_CATEGORY_INVALID", message="Invalid code."
                    )
                categories[row_index, index] = field.categories.index(str(value)) + 1
            else:
                stats = transform["continuous"][field.name]
                if not math.isfinite(stats["scale"]) or stats["scale"] <= 0:
                    raise DataContractError(code="CLINICAL_SCALE_INVALID", message="Invalid scale.")
                values[row_index, index] = (float(value) - stats["mean"]) / stats["scale"]
    result = BaselineClinical(values, categories, observed, schema_version)
    result.validate()
    return result
