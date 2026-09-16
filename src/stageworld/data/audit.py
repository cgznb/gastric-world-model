"""Redacted Excel structure audit that never emits cell sample values."""

from __future__ import annotations

import json
import os
import tempfile
from collections import Counter
from collections.abc import Mapping
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from stageworld.errors import DataContractError, StageWorldError


@dataclass(frozen=True, slots=True)
class ExcelAuditConfig:
    header_rows: tuple[int, ...] = (1,)
    sheet_name: str | None = None
    missing_tokens: frozenset[str] = frozenset(
        {"#n/a", "n/a", "na", "none", "null", "unknown", "missing"}
    )
    sensitive_header_patterns: tuple[str, ...] = (
        "name",
        "patient",
        "identifier",
        "identity",
        "phone",
        "address",
        "accession",
        "record number",
        "姓名",
        "患者",
        "身份证",
        "电话",
        "地址",
        "住院号",
        "摄片号",
        "病理号",
    )
    date_columns: frozenset[str] = frozenset()
    field_encodings: Mapping[str, frozenset[str]] = field(default_factory=dict)

    def __post_init__(self) -> None:
        if not self.header_rows or any(row < 1 for row in self.header_rows):
            raise DataContractError("invalid_header_rows", "header_rows must be positive")
        if tuple(sorted(set(self.header_rows))) != self.header_rows:
            raise DataContractError(
                "invalid_header_rows", "header_rows must be unique and ascending"
            )


@dataclass(frozen=True, slots=True)
class ColumnAudit:
    source_column_letter: str
    source_header_group: str | None
    source_header: str
    duplicate_header: bool
    nonmissing_count: int
    raw_missing_count: int
    semantic_missing_count: int
    unique_nonmissing_count: int
    type_counts: Mapping[str, int]
    formula_count: int
    error_count: int
    invalid_date_count: int
    unexpected_encoding_count: int


@dataclass(frozen=True, slots=True)
class ExcelAuditReport:
    source_sheet: str
    record_count: int
    column_count: int
    header_rows: tuple[int, ...]
    merged_range_count: int
    duplicate_header_count: int
    raw_missing_cell_count: int
    semantic_missing_cell_count: int
    formula_cell_count: int
    error_cell_count: int
    columns: tuple[ColumnAudit, ...]

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


def _require_openpyxl() -> Any:
    try:
        import openpyxl  # type: ignore[import-untyped]  # Optional clinical-I/O dependency.
    except ImportError as exc:
        raise StageWorldError(
            "dependency_missing",
            "Excel audit requires openpyxl",
            remediation="Install the clinical-io optional dependencies",
            details={"dependency": "openpyxl"},
        ) from exc
    return openpyxl


def _sanitize_header(value: object, patterns: tuple[str, ...]) -> str:
    text = "" if value is None else str(value).strip().replace("\n", " ")
    lowered = text.casefold()
    if any(pattern.casefold() in lowered for pattern in patterns):
        return "[REDACTED_IDENTIFIER_FIELD]"
    return text or "[UNNAMED_FIELD]"


def _semantic_missing(value: object, data_type: str, tokens: frozenset[str]) -> bool:
    if value is None or data_type == "e":
        return True
    if isinstance(value, str):
        return not value.strip() or value.strip().casefold() in tokens
    return False


def _value_kind(cell: Any) -> str:
    value = cell.value
    if cell.data_type == "f":
        return "formula"
    if cell.data_type == "e":
        return "error"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, (int, float)):
        return "numeric"
    if isinstance(value, str):
        return "string"
    return type(value).__name__


def _date_is_valid(value: object, epoch: object) -> bool:
    if isinstance(value, (datetime, date)):
        return True
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import (  # type: ignore[import-untyped]
                from_excel,
            )

            parsed = from_excel(value, epoch)
            return isinstance(parsed, (datetime, date))
        except (TypeError, ValueError, OverflowError):
            return False
    if isinstance(value, str):
        candidate = value.strip()
        for date_format in (
            "%Y-%m-%d",
            "%Y/%m/%d",
            "%Y.%m.%d",
            "%Y-%m-%d %H:%M:%S",
        ):
            try:
                datetime.strptime(candidate, date_format)
                return True
            except ValueError:
                continue
    return False


def audit_excel(path: str | Path, config: ExcelAuditConfig) -> ExcelAuditReport:
    openpyxl = _require_openpyxl()
    try:
        workbook = openpyxl.load_workbook(Path(path), read_only=False, data_only=False)
    except (OSError, ValueError) as exc:
        raise DataContractError("excel_unreadable", "Unable to read the Excel workbook") from exc
    if config.sheet_name is None:
        worksheet = workbook.active
    else:
        try:
            worksheet = workbook[config.sheet_name]
        except KeyError as exc:
            raise DataContractError(
                "excel_sheet_missing", "Configured worksheet is absent"
            ) from exc

    field_header_row = config.header_rows[-1]
    data_start_row = field_header_row + 1
    raw_headers = [
        worksheet.cell(field_header_row, column).value
        for column in range(1, worksheet.max_column + 1)
    ]
    normalized_headers = ["" if value is None else str(value).strip() for value in raw_headers]
    duplicate_headers = Counter(normalized_headers)

    groups: list[str | None] = [None] * worksheet.max_column
    if len(config.header_rows) > 1:
        group_rows = config.header_rows[:-1]
        previous_by_row = {row: "" for row in group_rows}
        for column in range(1, worksheet.max_column + 1):
            pieces = []
            for row in group_rows:
                value = worksheet.cell(row, column).value
                if value is not None and str(value).strip():
                    previous_by_row[row] = str(value).strip()
                if previous_by_row[row]:
                    pieces.append(previous_by_row[row])
            groups[column - 1] = " / ".join(dict.fromkeys(pieces)) or None

    columns: list[ColumnAudit] = []
    raw_missing_total = semantic_missing_total = formula_total = error_total = 0
    for column in range(1, worksheet.max_column + 1):
        letter = worksheet.cell(1, column).column_letter
        raw_missing = semantic_missing = formulas = errors = invalid_dates = unexpected = 0
        kinds: Counter[str] = Counter()
        unique: set[tuple[str, object]] = set()
        allowed_encoding = config.field_encodings.get(letter)
        for row in range(data_start_row, worksheet.max_row + 1):
            cell = worksheet.cell(row, column)
            value = cell.value
            is_raw_missing = value is None or (isinstance(value, str) and not value.strip())
            is_semantic_missing = _semantic_missing(value, cell.data_type, config.missing_tokens)
            raw_missing += int(is_raw_missing)
            semantic_missing += int(is_semantic_missing)
            formulas += int(cell.data_type == "f")
            errors += int(cell.data_type == "e")
            if is_semantic_missing:
                continue
            kinds[_value_kind(cell)] += 1
            try:
                unique.add((type(value).__name__, value))
            except TypeError:
                unique.add((type(value).__name__, repr(value)))
            if letter in config.date_columns and not _date_is_valid(value, workbook.epoch):
                invalid_dates += 1
            if allowed_encoding is not None and str(value).strip() not in allowed_encoding:
                unexpected += 1
        raw_missing_total += raw_missing
        semantic_missing_total += semantic_missing
        formula_total += formulas
        error_total += errors
        columns.append(
            ColumnAudit(
                source_column_letter=letter,
                source_header_group=(
                    _sanitize_header(groups[column - 1], config.sensitive_header_patterns)
                    if groups[column - 1] is not None
                    else None
                ),
                source_header=_sanitize_header(
                    raw_headers[column - 1], config.sensitive_header_patterns
                ),
                duplicate_header=duplicate_headers[normalized_headers[column - 1]] > 1,
                nonmissing_count=worksheet.max_row - field_header_row - semantic_missing,
                raw_missing_count=raw_missing,
                semantic_missing_count=semantic_missing,
                unique_nonmissing_count=len(unique),
                type_counts=dict(sorted(kinds.items())),
                formula_count=formulas,
                error_count=errors,
                invalid_date_count=invalid_dates,
                unexpected_encoding_count=unexpected,
            )
        )

    return ExcelAuditReport(
        source_sheet=_sanitize_header(worksheet.title, config.sensitive_header_patterns),
        record_count=max(worksheet.max_row - field_header_row, 0),
        column_count=worksheet.max_column,
        header_rows=config.header_rows,
        merged_range_count=len(worksheet.merged_cells.ranges),
        duplicate_header_count=sum(count - 1 for count in duplicate_headers.values() if count > 1),
        raw_missing_cell_count=raw_missing_total,
        semantic_missing_cell_count=semantic_missing_total,
        formula_cell_count=formula_total,
        error_cell_count=error_total,
        columns=tuple(columns),
    )


def write_redacted_excel_audit(report: ExcelAuditReport, path: str | Path) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    content = json.dumps(report.as_dict(), indent=2, sort_keys=True, ensure_ascii=True) + "\n"
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w", encoding="ascii", dir=target.parent, prefix=f".{target.name}.", delete=False
        ) as handle:
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())
            temporary = Path(handle.name)
        os.replace(temporary, target)
    except Exception:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        raise
