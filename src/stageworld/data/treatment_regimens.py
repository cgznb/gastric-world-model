"""Allowlisted interval-summary descriptors, never raw clinical text or dosing dates."""

from __future__ import annotations

import math
import re
import statistics
import unicodedata
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import replace
from pathlib import Path
from typing import Any

import torch

from stageworld.errors import ArtifactError, DataContractError
from stageworld.model import ActionTokens
from stageworld.training import WorldModelBatch

REGIMEN_PROTOCOL = "gastric-roi-regimen-os-development-v1"
REGIMEN_100EP_PROTOCOL = "gastric-roi-regimen-os-100ep-development-v1"
TUMOR_REGIMEN_100EP_PROTOCOL = "flare23-gastric-candidate-regimen-os-100ep-development-v1"
COARSE_TUMOR_100EP_PROTOCOL = "flare23-tumor-coarse-fallback-regimen-os-100ep-development-v1"
REGIMEN_FULL_EPOCH_PROTOCOLS = (
    REGIMEN_100EP_PROTOCOL, TUMOR_REGIMEN_100EP_PROTOCOL, COARSE_TUMOR_100EP_PROTOCOL,
)
REGIMEN_PROTOCOLS = (REGIMEN_PROTOCOL, *REGIMEN_FULL_EPOCH_PROTOCOLS)
REGIMEN_SCHEMA = "paired-ct-named-regimen-cycle-summary-v1"
REGIMEN_CODES = ("SOX", "XELOX", "CAPOX", "FOLFOX", "FLOT", "DOS", "DOX", "DCF", "FOLFIRI")
# Literal generic names only: do not infer ingredients, brands or therapeutic equivalence.
DRUG_NAMES = (
    ("oxaliplatin", "奥沙利铂"),
    ("tegafur_gimeracil_oteracil", "替吉奥"),
    ("capecitabine", "卡培他滨"),
    ("docetaxel", "多西他赛"),
    ("paclitaxel", "紫杉醇"),
    ("cisplatin", "顺铂"),
    ("irinotecan", "伊立替康"),
    ("fluorouracil", "氟尿嘧啶"),
    ("leucovorin", "亚叶酸钙"),
    ("sintilimab", "信迪利单抗"),
    ("camrelizumab", "卡瑞利珠单抗"),
    ("tislelizumab", "替雷利珠单抗"),
    ("toripalimab", "特瑞普利单抗"),
    ("nivolumab", "纳武利尤单抗"),
    ("pembrolizumab", "帕博利珠单抗"),
    ("serplulimab", "斯鲁利单抗"),
    ("penpulimab", "派安普利单抗"),
    ("durvalumab", "度伐利尤单抗"),
    ("atezolizumab", "阿替利珠单抗"),
    ("trastuzumab", "曲妥珠单抗"),
    ("bevacizumab", "贝伐珠单抗"),
    ("apatinib", "阿帕替尼"),
    ("anlotinib", "安罗替尼"),
    ("ramucirumab", "雷莫西尤单抗"),
    ("zolbetuximab", "佐妥昔单抗"),
)
MENTION_NAMES = tuple(f"regimen_{code.lower()}" for code in REGIMEN_CODES) + tuple(
    name for name, _ in DRUG_NAMES
)
CYCLE_FIELDS = (
    ("systemic", "AH", "系统治疗周期"),
    ("chemotherapy", "AI", "化疗周期"),
    ("immunotherapy", "AJ", "免疫治疗周期"),
)
DESCRIPTOR_COUNT = 5 + len(MENTION_NAMES) + len(CYCLE_FIELDS)
REGIMEN_ACTION_DIM = DESCRIPTOR_COUNT + 3
GROUPED_COLUMNS = ("AE", "AF", "AG", "AK", "AL")
MAX_SUMMARY_CYCLES = 60
_UNRESOLVED_CONTEXT = re.compile(
    r"未用|未予|未接受|未行|不使用|不用|否认|无需|计划|建议|拟|考虑|可能|[?？]"
)


def parse_named_mentions(value: Any, data_type: str) -> dict[str, int | None]:
    result: dict[str, int | None] = dict.fromkeys(MENTION_NAMES)
    if not isinstance(value, str) or data_type in {"e", "f"}:
        return result
    text = unicodedata.normalize("NFKC", value).upper()
    # Fail conservatively for negated/planned/uncertain summaries; no clinical NLP inference.
    if _UNRESOLVED_CONTEXT.search(text):
        return result
    for code in REGIMEN_CODES:
        if re.search(r"(?<![A-Z])" + code + r"(?![A-Z])", text):
            result[f"regimen_{code.lower()}"] = 1
    for name, literal in DRUG_NAMES:
        if literal in text:
            result[name] = 1
    return result


def parse_cycle_count(value: Any, data_type: str) -> tuple[int | None, str]:
    if data_type in {"e", "f"}:
        return None, "error_or_formula"
    if value is None:
        return None, "missing"
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        return None, "invalid"
    try:
        number = float(value)
    except (ValueError, OverflowError):
        return None, "invalid"
    if not math.isfinite(number) or not number.is_integer() or number < 0:
        return None, "invalid"
    if number > MAX_SUMMARY_CYCLES:
        return None, "outside_protocol_range"
    return int(number), "observed"


def read_regimen_rows(
    workbook_path: Path, *, patient_ids: set[str], pseudonymizer: Any
) -> list[dict[str, Any]]:
    import openpyxl  # type: ignore[import-untyped]
    from openpyxl.utils import column_index_from_string as ci  # type: ignore[import-untyped]

    from stageworld.data.paired_ct import _identifier
    from stageworld.data.treatment_summary import TREATMENT_FIELDS, _header, parse_treatment_flag

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        group = next(sheet.iter_rows(min_row=1, max_row=1, max_col=38))
        headers = next(sheet.iter_rows(min_row=2, max_row=2, max_col=38))
        expected = {"A": "序列号", "AD": "术前治疗"}
        expected.update({col: header for _, col, header in CYCLE_FIELDS})
        expected.update(
            {col: field[3] for col, field in zip(GROUPED_COLUMNS, TREATMENT_FIELDS, strict=True)}
        )
        if _header(group[ci("AD") - 1].value) != _header(
            "阶段2——术前化疗方案、周期及CT评估"
        ) or any(_header(headers[ci(c) - 1].value) != _header(h) for c, h in expected.items()):
            raise DataContractError(
                code="REGIMEN_HEADER_MISMATCH", message="Bind the approved grouped workbook."
            )
        result: dict[str, dict[str, Any]] = {}
        for cells in sheet.iter_rows(min_row=3, max_col=38):
            key = _identifier(cells[0].value, cells[0].data_type)
            if key is None:
                continue
            patient = pseudonymizer.token("patient", key, prefix="P")
            if patient not in patient_ids:
                continue
            if patient in result:
                raise DataContractError(
                    code="TREATMENT_DUPLICATE_PATIENT", message="Duplicate treatment row."
                )
            methods = {}
            for col, (name, _, negative, _) in zip(GROUPED_COLUMNS, TREATMENT_FIELDS, strict=True):
                cell = cells[ci(col) - 1]
                methods[name] = parse_treatment_flag(cell.value, cell.data_type, negative)
            summary = cells[ci("AD") - 1]
            cycles, quality = {}, {}
            for name, col, _ in CYCLE_FIELDS:
                cell = cells[ci(col) - 1]
                cycles[name], quality[name] = parse_cycle_count(cell.value, cell.data_type)
                if name in methods and methods[name] == 0 and cycles[name] not in (None, 0):
                    cycles[name], quality[name] = None, "modality_count_conflict"
            result[patient] = {
                "patient_id": patient,
                "methods": methods,
                "named_mentions": parse_named_mentions(summary.value, summary.data_type),
                "cycles": cycles,
                "cycle_quality": quality,
            }
        if set(result) != patient_ids:
            raise DataContractError(
                code="TREATMENT_PATIENT_MISSING", message="A development treatment row is missing."
            )
        return [result[p] for p in sorted(result)]
    finally:
        workbook.close()


def summarize_regimens(rows: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    from stageworld.data.treatment_summary import summarize_treatments

    return {
        **summarize_treatments(rows),
        "drug_regimens_used": True,
        "cycle_counts_used": True,
        "source_schema": "grouped_workbook_row1_groups_row2_fields_stage2_AD_AL",
        "named_mention_counts": {
            name: sum(r["named_mentions"][name] == 1 for r in rows) for name in MENTION_NAMES
        },
        "patients_with_any_named_mention": sum(
            any(v == 1 for v in r["named_mentions"].values()) for r in rows
        ),
        "patients_with_named_regimen_code": sum(
            any(r["named_mentions"][f"regimen_{code.lower()}"] == 1 for code in REGIMEN_CODES)
            for r in rows
        ),
        "cycle_quality_counts": {
            name: dict(Counter(r["cycle_quality"][name] for r in rows))
            for name, _, _ in CYCLE_FIELDS
        },
        "missing_name_means": "not_documented_not_confirmed_absence",
        "text_interpretation": "literal_allowlist_mentions_not_adjudicated_regimen_reconstruction",
        "dose_or_cycle_dates_used": False,
        "raw_text_used_as_model_input": False,
        "descriptor_count": DESCRIPTOR_COUNT,
        "action_input_dim": REGIMEN_ACTION_DIM,
    }


def fit_cycle_transform(
    rows: Sequence[Mapping[str, Any]], training_ids: set[str]
) -> dict[str, Any]:
    if not training_ids or not training_ids <= {r["patient_id"] for r in rows}:
        raise DataContractError(
            code="CYCLE_FIT_PATIENTS_INVALID", message="Bind training patients."
        )
    result = {}
    for name, _, _ in CYCLE_FIELDS:
        values = [
            math.log1p(r["cycles"][name])
            for r in rows
            if r["patient_id"] in training_ids and r["cycles"][name] is not None
        ]
        mean = statistics.fmean(values) if values else 0.0
        scale = statistics.pstdev(values) if len(values) > 1 else 0.0
        result[name] = {
            "log_mean": mean,
            "log_scale": scale if scale > 1e-6 else 1.0,
            "n": len(values),
        }
    return result


def encode_regimens(
    batch: WorldModelBatch,
    rows: Mapping[str, Mapping[str, Any]],
    transform: Mapping[str, Any],
) -> WorldModelBatch:
    values = encode_regimen_values([rows[p] for p in batch.patient_ids], transform)
    actions = regimen_summary_actions(values, batch.ct1_acquisition_time,
                                      provenance="confirmed_pre_ct_interval_summary")
    return replace(batch, treatment_actions=actions)


def encode_regimen_values(
    rows: Sequence[Mapping[str, Any]], transform: Mapping[str, Any],
) -> torch.Tensor:
    """Encode interval descriptors independently of any CT1 observation or cohort."""
    from stageworld.data.treatment_summary import METHOD_NAMES

    values = torch.zeros(len(rows), DESCRIPTOR_COUNT, REGIMEN_ACTION_DIM)
    for i, row in enumerate(rows):
        if (
            set(row["methods"]) != set(METHOD_NAMES)
            or set(row["named_mentions"]) != set(MENTION_NAMES)
            or set(row["cycles"]) != {n for n, _, _ in CYCLE_FIELDS}
        ):
            raise ArtifactError(
                code="TREATMENT_FIELDS_CHANGED", message="Unexpected descriptor fields."
            )
        descriptors = [row["methods"][n] for n in METHOD_NAMES]
        descriptors += [row["named_mentions"][n] for n in MENTION_NAMES]
        for v in descriptors:
            if v is not None and (type(v) is not int or v not in (0, 1)):
                raise ArtifactError(code="TREATMENT_VALUE_INVALID", message="Invalid descriptor.")
        for name, _, _ in CYCLE_FIELDS:
            value = row["cycles"][name]
            if value is not None:
                if type(value) is not int or not 0 <= value <= MAX_SUMMARY_CYCLES:
                    raise ArtifactError(
                        code="TREATMENT_VALUE_INVALID", message="Invalid cycle count."
                    )
                value = (math.log1p(value) - transform[name]["log_mean"]) / transform[name][
                    "log_scale"
                ]
            descriptors.append(value)
        for j, value in enumerate(descriptors):
            values[i, j, j] = 1
            values[i, j, DESCRIPTOR_COUNT] = 0 if value is None else value
            values[i, j, DESCRIPTOR_COUNT + 1] = float(value is not None)
            values[i, j, DESCRIPTOR_COUNT + 2] = 1
    return values


def regimen_summary_actions(
    values: torch.Tensor, target_time: torch.Tensor, *, provenance: str,
) -> ActionTokens:
    times = target_time[:, None].expand(-1, DESCRIPTOR_COUNT).clone()
    shape = times.shape
    actions = ActionTokens(
        values=values,
        valid=torch.ones(shape, dtype=torch.bool, device=values.device),
        event_time=times,
        available_time=times.clone(),
        event_type=torch.full(shape, 5, dtype=torch.long, device=values.device),
        planned_or_delivered=torch.ones(shape, dtype=torch.long, device=values.device),
        provenance=(REGIMEN_SCHEMA, provenance, "not_dosing_times"),
    )
    actions.validate()
    return actions
