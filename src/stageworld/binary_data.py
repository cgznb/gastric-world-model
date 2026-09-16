"""Separate binary labels and patient folds from CT/clinical feature fitting."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

import numpy as np
import torch
from sklearn.model_selection import (  # type: ignore[import-untyped]
    StratifiedKFold,
    StratifiedShuffleSplit,
)

from stageworld.binary_endpoints import ENDPOINTS, LABEL_CONTRACT, BinaryEndpointBatch
from stageworld.config import StageWorldConfig
from stageworld.ct6_workflow import snapshot_once
from stageworld.data.baseline_clinical import (
    CT6_CLINICAL_SCHEMA,
    encode_baseline,
    fit_clinical_transform,
    read_baseline_rows,
)
from stageworld.data.paired_ct import HMACPseudonymizer, _identifier
from stageworld.data.treatment_compact import (
    compact_actions,
    encode_compact,
    fit_name_support,
    read_compact_rows,
)
from stageworld.encoders.base import ObservationTokens
from stageworld.errors import DataContractError
from stageworld.generated_training import GeneratedTrainingBatch
from stageworld.real_workflow import (
    RealFeatureBundle,
    _paired_batch,
    load_real_world_pretrain_batches,
)
from stageworld.synthetic_workflow import _slice_observation


def parse_binary_label(value: Any, *, positive: int, negative: int) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    text = str(value).strip()
    if text in (str(positive), f"{positive}.0"):
        return 1
    if text in (str(negative), f"{negative}.0"):
        return 0
    return None


def read_binary_labels(
    path: Path, patient_ids: set[str], pseudonymizer: HMACPseudonymizer
) -> dict[str, list[int | None]]:
    import openpyxl  # type: ignore[import-untyped]
    from openpyxl.utils import column_index_from_string as ci  # type: ignore[import-untyped]

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        headers = next(sheet.iter_rows(min_row=2, max_row=2, max_col=ci("CG"), values_only=True))
        if (
            str(headers[ci("BJ") - 1]).strip() != "pCR(1=\u662f\uff0c2=\u5426)"
            or str(headers[ci("CG") - 1]).strip() != "\u590d\u53d1\u8f6c\u79fb\u72b6\u6001"
        ):
            raise DataContractError(
                code="BINARY_HEADERS", message="Use the latest endpoint columns."
            )
        result: dict[str, list[int | None]] = {}
        for cells in sheet.iter_rows(min_row=3, max_col=ci("CG")):
            key = _identifier(cells[0].value, cells[0].data_type)
            if key is None:
                continue
            patient = pseudonymizer.token("patient", key, prefix="P")
            if patient not in patient_ids:
                continue
            if patient in result:
                raise DataContractError(
                    code="BINARY_DUPLICATE", message="Repeated endpoint patient."
                )
            result[patient] = []
            for name in ENDPOINTS:
                spec = LABEL_CONTRACT[name]
                cell = cells[ci(spec["column"]) - 1]
                result[patient].append(
                    None
                    if cell.data_type in ("f", "e")
                    else parse_binary_label(
                        cell.value, positive=spec["positive"], negative=spec["negative"]
                    )
                )
        if set(result) != patient_ids:
            raise DataContractError(
                code="BINARY_COVERAGE", message="Endpoint rows must cover the pool."
            )
        return result
    finally:
        workbook.close()


def label_counts(ids: list[str], labels: dict[str, list[int | None]]) -> dict[str, Any]:
    result: dict[str, Any] = {"patients": len(ids)}
    for index, name in enumerate(ENDPOINTS):
        values = [labels[p][index] for p in ids]
        result[name] = {
            "positive": values.count(1),
            "negative": values.count(0),
            "missing": values.count(None),
        }
    return result


def make_binary_folds(
    labels: dict[str, list[int | None]], *, seed: int = 17, folds: int = 5
) -> dict[str, Any]:
    ids = sorted(labels)
    y = np.array([[2 if x is None else x for x in labels[p]] for p in ids], dtype=int)
    strata = y[:, 0] + 3 * y[:, 1]
    if len(ids) < folds or min(Counter(strata).values()) < folds:
        raise DataContractError(
            code="BINARY_FOLD_SUPPORT", message="Joint label strata lack fold support."
        )
    assignment: dict[str, Any] = {
        "split_seed": seed,
        "fold_count": folds,
        "stratification": "joint_pcr_recurrence_missing_status",
        "inner_validation_fraction": 0.2,
        "pool": label_counts(ids, labels),
        "folds": [],
    }
    outer = StratifiedKFold(n_splits=folds, shuffle=True, random_state=seed)
    seen: set[str] = set()
    for fold, (pool, held) in enumerate(outer.split(np.zeros(len(ids)), strata)):
        inner = StratifiedShuffleSplit(n_splits=1, test_size=0.2, random_state=seed + fold + 1)
        train_idx, validation_idx = next(inner.split(np.zeros(len(pool)), strata[pool]))
        groups = {
            "train": [ids[i] for i in pool[train_idx]],
            "validation": [ids[i] for i in pool[validation_idx]],
            "outer": [ids[i] for i in held],
        }
        if any(
            set(groups[a]) & set(groups[b])
            for a, b in (("train", "validation"), ("train", "outer"), ("validation", "outer"))
        ):
            raise AssertionError("Overlapping fold patients")
        if seen & set(groups["outer"]):
            raise AssertionError("Repeated outer patient")
        seen.update(groups["outer"])
        assignment["folds"].append(
            {
                "fold": fold,
                "patient_ids": groups,
                "counts": {name: label_counts(group, labels) for name, group in groups.items()},
            }
        )
    if seen != set(ids):
        raise AssertionError("Outer-fold coverage differs")
    return assignment


@dataclass
class BinaryDevelopmentData:
    bundle: RealFeatureBundle
    features: dict[str, dict[str, ObservationTokens]]
    times: dict[tuple[str, str], float]
    clinical: dict[str, Any]
    treatments: dict[str, Any]
    labels: dict[str, list[int | None]]
    source_id: str


def prepare_binary_development(config: StageWorldConfig) -> BinaryDevelopmentData:
    bundle = load_real_world_pretrain_batches(config)
    ids = {
        p
        for batches in bundle.batches_by_split.values()
        for batch in batches
        for p in batch.patient_ids
    }
    if (
        len(ids) != 591
        or bundle.training_patient_count != 486
        or bundle.validation_patient_count != 105
    ):
        raise DataContractError(
            code="BINARY_POOL", message="Keep the original 591 development patients."
        )
    if not config.paths.clinical_excel or not config.paths.identity_hmac_key_file:
        raise DataContractError(
            code="BINARY_BINDINGS", message="Bind the latest clinical workbook."
        )
    path = Path(config.paths.clinical_excel)
    pseudonymizer = HMACPseudonymizer.from_file(
        Path(config.paths.identity_hmac_key_file), project_root=Path.cwd()
    )
    clinical, audit = read_baseline_rows(
        path, ids, pseudonymizer, schema_version=CT6_CLINICAL_SCHEMA
    )
    treatments = read_compact_rows(path, ids, pseudonymizer)
    labels = read_binary_labels(path, ids, pseudonymizer)
    snapshot = snapshot_once(
        config.output_root / "raw_inputs_and_labels.json",
        {
            "clinical_rows": clinical,
            "clinical_audit": audit,
            "treatment_rows": treatments,
            "labels": labels,
            "label_contract": LABEL_CONTRACT,
            "feature_artifact_id": bundle.feature_artifact_id,
            "cohort_artifact_id": bundle.cohort_artifact_id,
            "test_used": False,
        },
        "binary-source",
    )
    features: dict[str, dict[str, ObservationTokens]] = {}
    times: dict[tuple[str, str], float] = {}
    for batches in bundle.batches_by_split.values():
        for batch in batches:
            for i, patient in enumerate(batch.patient_ids):
                index = torch.tensor([i])
                features[patient] = {
                    "baseline_ct": _slice_observation(batch.ct0, index),
                    "post_treatment_ct": _slice_observation(batch.ct1, index),
                }
                times[patient, "s0"] = float(batch.s0_time[i])
                times[patient, "s1"] = float(batch.s1_time[i])
    return BinaryDevelopmentData(
        bundle,
        features,
        times,
        clinical,
        {row["patient_id"]: row for row in treatments},
        labels,
        snapshot["artifact_id"],
    )


def make_fold_data(
    config: StageWorldConfig,
    data: BinaryDevelopmentData,
    fold: dict[str, Any],
    root: Path,
    *,
    smoke: bool = False,
) -> tuple[RealFeatureBundle, tuple[BinaryEndpointBatch, ...], dict[str, Any]]:
    groups = {k: (v[:32] if smoke else v) for k, v in fold["patient_ids"].items()}
    train_ids = set(groups["train"])
    transform = fit_clinical_transform(data.clinical, train_ids, schema_version=CT6_CLINICAL_SCHEMA)
    support = fit_name_support(list(data.treatments.values()), train_ids)
    snapshot = snapshot_once(
        root / "inputs.json",
        {
            "source_id": data.source_id,
            "fold": fold["fold"],
            "smoke": smoke,
            "split_patient_ids": groups,
            "transform": transform,
            "treatment_support": support,
            "clinical_schema": CT6_CLINICAL_SCHEMA,
            "label_contract": LABEL_CONTRACT,
            "fit_patient_ids": sorted(train_ids),
        },
        "binary-fold-inputs",
    )
    batches = {}
    for split, ids in groups.items():
        converted = []
        for start in range(0, len(ids), config.training.patient_batch_size):
            selected = ids[start : start + config.training.patient_batch_size]
            base = _paired_batch(config, selected, data.features, data.times)
            values, _ = encode_compact([data.treatments[p] for p in selected], support)
            base = replace(
                base,
                treatment_actions=compact_actions(
                    values,
                    base.ct1_acquisition_time,
                    provenance="retrospective_observed_interval_condition_not_baseline_fact",
                ),
            )
            generated = GeneratedTrainingBatch.from_world(
                base, encode_baseline([data.clinical[p] for p in selected], transform)
            )
            valid = torch.tensor([[x is not None for x in data.labels[p]] for p in selected])
            labels = torch.tensor(
                [[0 if x is None else x for x in data.labels[p]] for p in selected],
                dtype=torch.float32,
            )
            converted.append(BinaryEndpointBatch.from_generated(generated, labels, valid))
        batches[split] = tuple(converted)
    return (
        replace(
            data.bundle,
            batches_by_split={k: batches[k] for k in ("train", "validation")},
            training_patient_count=len(groups["train"]),
            validation_patient_count=len(groups["validation"]),
            split_version=snapshot["artifact_id"],
            treatment_snapshot=snapshot,
        ),
        batches["outer"],
        snapshot,
    )


def unlabelled_bundle(bundle: RealFeatureBundle) -> RealFeatureBundle:
    if any(
        not isinstance(batch, BinaryEndpointBatch)
        for batches in bundle.batches_by_split.values()
        for batch in batches
    ):
        raise DataContractError(
            code="BINARY_BATCH", message="Bind endpoint batches before masking."
        )
    return replace(
        bundle,
        batches_by_split={
            name: tuple(
                replace(
                    batch,
                    endpoint_labels=torch.zeros_like(batch.endpoint_labels),
                    endpoint_valid=torch.zeros_like(batch.endpoint_valid),
                )
                for batch in batches
                if isinstance(batch, BinaryEndpointBatch)
            )
            for name, batches in bundle.batches_by_split.items()
        },
    )
