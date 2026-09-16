"""Local host binding without exposing clinical paths or credentials in logs."""

from __future__ import annotations

import argparse
import contextlib
import json
import os
import sys
import traceback
from dataclasses import replace
from pathlib import Path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "phase",
        choices=("extract", "train", "develop", "verify", "audit-treatments", "prepare-treatments"),
    )
    parser.add_argument("--limit-per-split", type=int)
    parser.add_argument("--with-treatments", action="store_true")
    parser.add_argument("--with-regimens", action="store_true")
    parser.add_argument("--epochs-100", action="store_true")
    roi_mode = parser.add_mutually_exclusive_group()
    roi_mode.add_argument("--tumor-candidates", action="store_true")
    roi_mode.add_argument("--coarse-tumor-roi", action="store_true")
    parser.add_argument("--verify-after-training", action="store_true")
    args = parser.parse_args()
    tumor_mode = args.tumor_candidates or args.coarse_tumor_roi
    if tumor_mode:
        args.with_regimens = True
        if args.phase == "develop":
            parser.error("Tumor candidates require extract, prepare-treatments, then train.")
    if args.epochs_100 and (not args.with_regimens or args.phase not in ("train", "verify")):
        parser.error("--epochs-100 requires --with-regimens and train or verify.")
    if args.verify_after_training and args.phase not in ("train", "develop"):
        parser.error("--verify-after-training requires train or develop.")
    project = Path(__file__).resolve().parents[1]
    os.chdir(project)
    os.umask(0o077)
    from stageworld.artifacts import atomic_write_json, read_json
    from stageworld.config import load_config
    from stageworld.data.gastric_roi import offline_network
    from stageworld.data.paired_ct import HMACPseudonymizer
    from stageworld.real_workflow import extract_real_ct_features

    assets = read_json(
        project
        / "artifacts/real/paired-ct-os-v1-first-acquisition"
        / "data/restricted/asset_bindings.json"
    )
    clinical_root = os.path.commonpath([row["local_path"] for row in assets["bindings"]])
    # Regimen runs read the explicitly supplied workbook and reuse frozen CT caches.
    workbook_binding = os.environ.get("STAGEWORLD_CLINICAL_EXCEL")
    approved_input_root = (
        str(Path(workbook_binding).expanduser().resolve().parent)
        if args.with_regimens and workbook_binding and args.phase != "extract"
        else clinical_root
    )
    os.environ.setdefault("STAGEWORLD_WEIAI_ROOT", approved_input_root)
    for variable, filename in (
        ("STAGEWORLD_SWINUNETR_WEIGHT", "model_swinvit.pt"),
        ("STAGEWORLD_HMAC_KEY_FILE", "gastric-os-v1.hmac"),
    ):
        if variable not in os.environ:
            matches = list(Path.home().rglob(filename))
            if len(matches) != 1:
                print(json.dumps({"status": "blocked", "missing_binding": variable}))
                return 2
            os.environ[variable] = str(matches[0])
    os.environ.setdefault(
        "STAGEWORLD_GASTRIC_ROI_FEATURE_ROOT",
        str(Path.home() / ".local/share/stageworld/gastric-roi-os-resident-v1"),
    )
    os.environ.setdefault(
        "STAGEWORLD_TOTALSEG_HOME", str(Path.home() / ".cache/stageworld-models/totalseg-2.18.0")
    )
    if tumor_mode:
        os.environ.setdefault(
            "STAGEWORLD_TUMOR_MODEL",
            str(Path.home() / ".cache/stageworld-models/flare23-blackbean-20260911"),
        )
        os.environ.setdefault(
            "STAGEWORLD_TUMOR_ROI_FEATURE_ROOT",
            str(Path.home() / (
                ".local/share/stageworld/flare23-tumor-coarse-fallback-os-v1"
                if args.coarse_tumor_roi else
                ".local/share/stageworld/flare23-gastric-candidate-os-v1"
            )),
        )
        os.environ.setdefault(
            "STAGEWORLD_TUMOR_TREATMENT_MANIFEST",
            str(Path(os.environ["STAGEWORLD_TUMOR_ROI_FEATURE_ROOT"]) / "treatment_manifest.json"),
        )
    os.environ.setdefault(
        "STAGEWORLD_TREATMENT_MANIFEST",
        str(
            Path.home()
            / (
                ".local/share/stageworld/gastric-regimen-os-v1/treatment_manifest.json"
                if args.with_regimens
                else ".local/share/stageworld/gastric-treatment-os-v1/treatment_manifest.json"
            )
        ),
    )
    import torch

    torch.set_num_threads(4)
    torch.set_num_interop_threads(2)
    with_treatments = (
        args.with_regimens or args.with_treatments or args.phase.endswith("treatments")
    )
    config = load_config(
        project
        / "configs"
        / (
            "project.weiai-os-v1-tumor-coarse-roi-100ep.yaml"
            if args.coarse_tumor_roi else "project.weiai-os-v1-tumor-candidate-roi-100ep.yaml"
            if args.tumor_candidates
            else "project.weiai-os-v1-regimen-roi-100ep.yaml"
            if args.epochs_100
            else "project.weiai-os-v1-regimen-roi.yaml"
            if args.with_regimens
            else "project.weiai-os-v1-treatment-roi.yaml"
            if with_treatments
            else "project.weiai-os-v1-gastric-roi.yaml"
        )
    )
    config.validate(command=args.phase, supervised=args.phase != "extract")
    log_root = config.output_root / "restricted_logs"
    log_root.mkdir(parents=True, exist_ok=True, mode=0o700)
    # Native image libraries may emit sensitive header/path diagnostics to file descriptors.
    with (log_root / f"{args.phase}.log").open("a") as log:
        saved_out, saved_err = os.dup(1), os.dup(2)
        os.dup2(log.fileno(), 1)
        os.dup2(log.fileno(), 2)
        try:
            with (
                contextlib.redirect_stdout(log),
                contextlib.redirect_stderr(log),
                offline_network(),
            ):
                if (
                    with_treatments
                    and not tumor_mode
                    and args.phase in ("extract", "develop")
                ):
                    raise ValueError("Treatment runs reuse the fixed CT cache; use train.")
                if tumor_mode and args.phase in ("train", "verify"):
                    from stageworld.real_workflow import load_real_world_pretrain_batches

                    bundle = load_real_world_pretrain_batches(config)
                    steps = 100 * len(bundle.batches_by_split["train"])
                    config = replace(
                        config,
                        training=replace(
                            config.training, world_pretrain_steps=steps, joint_survival_steps=steps
                        ),
                    )
                if args.phase in ("audit-treatments", "prepare-treatments"):
                    from stageworld.data.treatment_summary import prepare_treatments
                    from stageworld.real_workflow import load_real_world_pretrain_batches

                    if config.paths.clinical_excel is None:
                        if args.with_regimens:
                            raise ValueError("Explicitly bind the approved grouped workbook.")
                        import openpyxl

                        candidates = []
                        search_root = Path(
                            os.environ.get(
                                "STAGEWORLD_WORKBOOK_SEARCH_ROOT", str(Path(clinical_root).parent)
                            )
                        )
                        for pattern in ("*.xlsx", "*/*.xlsx", "*/*/*.xlsx", "*/*/*/*.xlsx"):
                            for candidate in search_root.glob(pattern):
                                with contextlib.suppress(Exception):
                                    workbook = openpyxl.load_workbook(
                                        candidate, read_only=True, data_only=False
                                    )
                                    try:
                                        sheet = workbook.worksheets[0]
                                        if sheet.max_column == 90 and sheet.max_row == 957:
                                            candidates.append(candidate)
                                    finally:
                                        workbook.close()
                        if len(candidates) != 1:
                            raise ValueError("Bind exactly one approved clinical workbook.")
                        config = replace(
                            config, paths=replace(config.paths, clinical_excel=str(candidates[0]))
                        )
                    result = prepare_treatments(
                        config,
                        load_real_world_pretrain_batches(config),
                        HMACPseudonymizer.from_file(
                            Path(os.environ["STAGEWORLD_HMAC_KEY_FILE"]), project_root=project
                        ),
                        audit_only=args.phase == "audit-treatments",
                    )
                    atomic_write_json(config.output_root / "treatment_audit.json", result)
                if args.phase in ("extract", "develop"):
                    if args.coarse_tumor_roi:
                        from stageworld.data.tumor_roi import reuse_tumor_mask_predictions

                        pilot_masks = Path.home() / (
                            ".local/share/stageworld/flare23-gastric-pilot-20260911/tumor_masks"
                        )
                        if pilot_masks.is_dir():
                            target_masks = Path(os.environ["STAGEWORLD_TUMOR_ROI_FEATURE_ROOT"])
                            reused = reuse_tumor_mask_predictions(
                                Path(os.environ["STAGEWORLD_TUMOR_MODEL"]), pilot_masks,
                                target_masks / "tumor_masks",
                            )
                            atomic_write_json(config.output_root / "mask_reuse.json", reused)
                    key_file = Path(os.environ["STAGEWORLD_HMAC_KEY_FILE"])
                    result = extract_real_ct_features(
                        config,
                        pseudonymizer=HMACPseudonymizer.from_file(key_file, project_root=project),
                        limit_per_split=args.limit_per_split,
                        include_test=False,
                        device="cuda",
                    )
                if args.phase in ("train", "develop"):
                    from stageworld.real_survival import run_real_os_development

                    result = run_real_os_development(config)
                    if args.verify_after_training:
                        from stageworld.real_survival import verify_real_os_development

                        result["verification"] = verify_real_os_development(config)
                if args.phase == "verify":
                    from stageworld.real_survival import verify_real_os_development

                    result = verify_real_os_development(config)
            status = 0
        except BaseException as error:
            traceback.print_exc(file=log)
            result = {
                "status": "failed",
                "error_code": getattr(error, "code", type(error).__name__),
            }
            status = 1
        finally:
            log.flush()
            os.dup2(saved_out, 1)
            os.dup2(saved_err, 2)
            os.close(saved_out)
            os.close(saved_err)
    print(json.dumps(result, ensure_ascii=True))
    return status


if __name__ == "__main__":
    sys.exit(main())
