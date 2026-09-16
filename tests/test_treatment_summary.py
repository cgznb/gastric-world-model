from dataclasses import replace
from pathlib import Path

import pytest
import torch
from test_training import _batch, _model

from stageworld.artifacts import atomic_write_private_json, read_json
from stageworld.config import load_config
from stageworld.data.paired_ct import HMACPseudonymizer
from stageworld.data.treatment_summary import (
    CONFIRMED_CUTOFF,
    METHOD_NAMES,
    TREATMENT_FIELDS,
    TREATMENT_SCHEMA,
    attach_treatments,
    encode_treatments,
    parse_treatment_flag,
    read_treatment_rows,
    require_treatment_cutoff,
    summarize_treatments,
)
from stageworld.errors import ArtifactError, DataContractError
from stageworld.model import ActionTokens, StageWorldModel
from stageworld.real_survival import _baseline_inputs, _forward, treatment_response_probe


def treatment_batch(indices, action_input_dim=8):
    batch = _batch(indices)
    empty = ActionTokens.empty(
        batch_size=batch.batch_size, value_dim=action_input_dim, device=torch.device("cpu")
    )
    return replace(batch, treatment_actions=empty, surgery_actions=empty)


def write_treatment_fixture(config, bundle):
    rows = [
        {"patient_id": p, "methods": {name: (i + j) % 2 for j, name in enumerate(METHOD_NAMES)}}
        for i, p in enumerate(
            p
            for batches in bundle.batches_by_split.values()
            for b in batches
            for p in b.patient_ids
        )
    ]
    payload = {
        "schema_version": TREATMENT_SCHEMA,
        "artifact_id": "treatment-interval-synthetic-fixture",
        "data_lineage_id": bundle.data_lineage_id,
        "cohort_artifact_id": bundle.cohort_artifact_id,
        "split_version": bundle.split_version,
        "ct_feature_artifact_id": bundle.feature_artifact_id,
        "cutoff": CONFIRMED_CUTOFF,
        "rows": rows,
        "aggregate": summarize_treatments(rows),
    }
    path = Path(config.paths.treatment_manifest)
    atomic_write_private_json(path.parent / "versions" / f"{payload['artifact_id']}.json", payload)
    atomic_write_private_json(path, payload)
    return payload


@pytest.mark.parametrize("negative", [0, 2])
def test_treatment_codes_and_unknowns_are_not_interchangeable(negative):
    assert parse_treatment_flag(1, "n", negative) == 1
    assert parse_treatment_flag(str(negative), "s", negative) == 0
    assert parse_treatment_flag(2 if negative == 0 else 0, "n", negative) is None
    for value, kind in [
        (None, "n"),
        ("#N/A", "e"),
        ("=1", "f"),
        (True, "b"),
        (float("nan"), "n"),
        ("free text", "s"),
    ]:
        assert parse_treatment_flag(value, kind, negative) is None


def test_workbook_treatments_use_exact_columns_and_skip_other_patients(tmp_path):
    import openpyxl
    from openpyxl.utils import column_index_from_string

    pseudonymizer = HMACPseudonymizer(bytes(range(32)))
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for _, column, _, header in TREATMENT_FIELDS:
        sheet.cell(1, column_index_from_string(column), header)
    for row in range(2, 5):
        sheet.cell(row, 1, f"row-{row}")
        for _, column, negative, _ in TREATMENT_FIELDS:
            sheet.cell(row, column_index_from_string(column), 1 if row == 2 else negative)
        sheet.cell(row, 71, "forbidden-outcome-text")
    sheet.cell(3, 28, 0)  # AB documents 2, not 0, as negative.
    path = tmp_path / "synthetic.xlsx"
    workbook.save(path)
    ids = {pseudonymizer.token("patient", f"row-{r}", prefix="P") for r in (2, 3)}
    rows = read_treatment_rows(path, patient_ids=ids, pseudonymizer=pseudonymizer)
    assert {r["patient_id"] for r in rows} == ids
    aggregate = summarize_treatments(rows)
    assert aggregate["modalities"]["chemotherapy"] == {"present": 1, "absent": 1}
    assert aggregate["modalities"]["immunotherapy"] == {"present": 1, "unknown": 1}
    assert "forbidden-outcome-text" not in str(rows)
    sheet.cell(1, 28, "unexpected-header")
    workbook.save(path)
    with pytest.raises(DataContractError) as error:
        read_treatment_rows(path, patient_ids=ids, pseudonymizer=pseudonymizer)
    assert error.value.code == "TREATMENT_HEADER_MISMATCH"


def test_unconfirmed_interval_is_blocked_before_training():
    config = load_config(
        Path(__file__).resolve().parents[1] / "configs/project.weiai-os-v1-treatment-roi.yaml"
    )
    with pytest.raises(DataContractError) as error:
        require_treatment_cutoff(
            replace(config, clinical=replace(config.clinical, treatment_summary_cutoff=None))
        )
    assert error.value.code == "TREATMENT_CT_CUTOFF_UNCONFIRMED"


def test_treatment_affects_prior_and_s1_with_gradients_but_not_s0():
    batch = treatment_batch((0, 1, 2, 3))
    rows = {p: {"methods": {name: 0 for name in METHOD_NAMES}} for p in batch.patient_ids}
    batch = encode_treatments(batch, rows)
    assert torch.equal(batch.treatment_actions.event_time[:, 0], batch.ct1_acquisition_time)
    model = StageWorldModel(replace(_model().config, action_input_dim=8))
    probe = treatment_response_probe(model, (batch,))
    assert probe["maximum_absolute_differences"]["s0_rate"] == 0
    assert all(
        probe["maximum_absolute_differences"][key] > 0
        for key in ("prior_state", "future_ct", "s1_rate")
    )
    actions = replace(
        batch.treatment_actions, values=batch.treatment_actions.values.clone().requires_grad_()
    )
    output = _forward(model, replace(batch, treatment_actions=actions))
    assert (
        torch.autograd.grad(
            output.survival_s0.rates.sum(), actions.values, allow_unused=True, retain_graph=True
        )[0]
        is None
    )
    gradient = torch.autograd.grad(output.future_ct.mean.square().sum(), actions.values)[0]
    assert torch.isfinite(gradient).all() and gradient[:, :, 5].abs().sum() > 0
    changed = replace(
        batch,
        treatment_actions=replace(
            batch.treatment_actions, values=batch.treatment_actions.values + 1
        ),
    )
    assert torch.equal(_baseline_inputs((batch,), 0), _baseline_inputs((changed,), 0))
    assert not torch.equal(_baseline_inputs((batch,), 1), _baseline_inputs((changed,), 1))


def test_treatment_manifest_has_exact_snapshot_and_population(tmp_path):
    from stageworld.cache import CacheProvenance
    from stageworld.real_workflow import RealFeatureBundle

    config = load_config(
        Path(__file__).resolve().parents[1] / "configs/project.weiai-os-v1-treatment-roi.yaml"
    )
    config = replace(
        config,
        paths=replace(config.paths, treatment_manifest=str(tmp_path / "treatments/manifest.json")),
        clinical=replace(config.clinical, treatment_summary_cutoff=CONFIRMED_CUTOFF),
    )
    batch = treatment_batch((0, 1))
    provenance = CacheProvenance.from_encoder(
        batch.ct0.provenance,
        schema_version="test",
        patch_sampling_version="test",
        split_version="test",
        target_transform_version="test",
        teacher_version="test",
    )
    bundle = RealFeatureBundle(
        batches_by_split={"train": (batch,)},
        data_lineage_id="test",
        cohort_artifact_id="test",
        split_version="test",
        feature_artifact_id="test",
        cache_provenance=provenance,
        training_patient_count=2,
        validation_patient_count=0,
        held_out_test_patient_count=1,
    )
    payload = write_treatment_fixture(config, bundle)
    attached = attach_treatments(config, bundle)
    assert attached.treatment_snapshot == payload
    assert attached.batches_by_split["train"][0].ct0 is batch.ct0
    path = Path(config.paths.treatment_manifest)
    bad = read_json(path)
    bad["rows"][0]["methods"]["chemotherapy"] = 1 - bad["rows"][0]["methods"]["chemotherapy"]
    atomic_write_private_json(path, bad)
    with pytest.raises(ArtifactError) as error:
        attach_treatments(config, bundle)
    assert error.value.code == "TREATMENT_SNAPSHOT_MISMATCH"
