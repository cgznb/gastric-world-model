from dataclasses import replace
from pathlib import Path

import pytest
import torch
from test_training import _model
from test_treatment_summary import treatment_batch

from stageworld.artifacts import atomic_write_private_json
from stageworld.data.paired_ct import HMACPseudonymizer
from stageworld.data.treatment_regimens import (
    CYCLE_FIELDS,
    DESCRIPTOR_COUNT,
    GROUPED_COLUMNS,
    MENTION_NAMES,
    REGIMEN_ACTION_DIM,
    REGIMEN_SCHEMA,
    encode_regimens,
    fit_cycle_transform,
    parse_cycle_count,
    parse_named_mentions,
    read_regimen_rows,
    summarize_regimens,
)
from stageworld.data.treatment_summary import CONFIRMED_CUTOFF, METHOD_NAMES, TREATMENT_FIELDS
from stageworld.errors import DataContractError
from stageworld.model import StageWorldModel
from stageworld.real_survival import _baseline_inputs, _forward, treatment_response_probe


def regimen_rows(patients):
    return [
        {
            "patient_id": patient,
            "methods": {name: i % 2 for name in METHOD_NAMES},
            "named_mentions": parse_named_mentions("SOX" if i % 2 else "XELOX", "s"),
            "cycles": {name: 2 + i % 4 for name, _, _ in CYCLE_FIELDS},
            "cycle_quality": {name: "observed" for name, _, _ in CYCLE_FIELDS},
        }
        for i, patient in enumerate(patients)
    ]


def write_regimen_fixture(config, bundle):
    rows = regimen_rows(
        p for batches in bundle.batches_by_split.values() for b in batches for p in b.patient_ids
    )
    training_ids = {p for b in bundle.batches_by_split["train"] for p in b.patient_ids}
    payload = {
        "schema_version": REGIMEN_SCHEMA,
        "artifact_id": "treatment-interval-synthetic-fixture",
        "data_lineage_id": bundle.data_lineage_id,
        "cohort_artifact_id": bundle.cohort_artifact_id,
        "split_version": bundle.split_version,
        "ct_feature_artifact_id": bundle.feature_artifact_id,
        "cutoff": CONFIRMED_CUTOFF,
        "rows": rows,
        "aggregate": summarize_regimens(rows),
        "cycle_transform": fit_cycle_transform(rows, training_ids),
        "cycle_transform_fit_split": "train",
    }
    path = Path(config.paths.treatment_manifest)
    atomic_write_private_json(path.parent / "versions" / f"{payload['artifact_id']}.json", payload)
    atomic_write_private_json(path, payload)
    return payload


def test_regimen_mentions_are_allowlisted_not_raw_text_or_invented_ingredients():
    mentions = parse_named_mentions("SOX4 + XELOX2 + 信迪利单抗; PRIVATE-NOTE-DO-NOT-RETAIN", "s")
    assert mentions["regimen_sox"] == mentions["regimen_xelox"] == mentions["sintilimab"] == 1
    assert mentions["oxaliplatin"] is None
    assert set(mentions) == set(MENTION_NAMES)
    assert "PRIVATE-NOTE" not in str(mentions)
    assert parse_named_mentions("unknown scheme", "s") == dict.fromkeys(MENTION_NAMES)
    assert parse_named_mentions("NOTSOX and SOXYZ", "s")["regimen_sox"] is None
    for value in ("未用SOX", "计划XELOX", "考虑FLOT", "SOX?"):
        assert all(v is None for v in parse_named_mentions(value, "s").values())


@pytest.mark.parametrize(
    "value,kind",
    [
        (None, "n"),
        ("#N/A", "e"),
        ("=4", "f"),
        (True, "b"),
        (-1, "n"),
        (2.5, "n"),
        (61, "n"),
        (float("inf"), "n"),
    ],
)
def test_invalid_cycles_stay_unknown(value, kind):
    assert parse_cycle_count(value, kind)[0] is None
    assert parse_cycle_count(0, "n") == (0, "observed")
    assert parse_cycle_count("4", "s") == (4, "observed")


def test_cycle_transform_is_training_only_and_regimens_change_transition_not_s0():
    batch = treatment_batch((0, 1, 2, 3), REGIMEN_ACTION_DIM)
    rows = regimen_rows(batch.patient_ids)
    training_ids = set(batch.patient_ids[:2])
    transform = fit_cycle_transform(rows, training_ids)
    rows[-1]["cycles"]["systemic"] = 60
    assert fit_cycle_transform(rows, training_ids) == transform
    batch = encode_regimens(batch, {r["patient_id"]: r for r in rows}, transform)
    assert batch.treatment_actions.values.shape == (4, DESCRIPTOR_COUNT, REGIMEN_ACTION_DIM)
    model = StageWorldModel(replace(_model().config, action_input_dim=REGIMEN_ACTION_DIM))
    probe = treatment_response_probe(model, (batch,), regimen=True)
    assert probe["s0_invariant"] is True
    assert all(
        probe["maximum_absolute_differences"][key] > 0
        for key in ("prior_state", "future_ct", "s1_rate")
    )
    actions = replace(
        batch.treatment_actions, values=batch.treatment_actions.values.clone().requires_grad_()
    )
    changed = replace(batch, treatment_actions=actions)
    output = _forward(model, changed)
    assert (
        torch.autograd.grad(
            output.survival_s0.rates.sum(), actions.values, allow_unused=True, retain_graph=True
        )[0]
        is None
    )
    gradient = torch.autograd.grad(output.future_ct.mean.square().sum(), actions.values)[0]
    assert gradient[:, 5:, DESCRIPTOR_COUNT].abs().sum() > 0
    perturbed = replace(batch, treatment_actions=replace(actions, values=actions.values + 1))
    assert torch.equal(_baseline_inputs((batch,), 0), _baseline_inputs((perturbed,), 0))
    assert not torch.equal(_baseline_inputs((batch,), 1), _baseline_inputs((perturbed,), 1))


def test_grouped_reader_uses_row2_headers_ignores_outcomes_and_test_rows(tmp_path):
    import openpyxl
    from openpyxl.utils import column_index_from_string as ci

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["AD1"] = "阶段2——术前化疗方案、周期及CT评估"
    sheet["A2"] = "序列号"
    sheet["AD2"] = "术前治疗"
    for column, field in zip(GROUPED_COLUMNS, TREATMENT_FIELDS, strict=True):
        sheet.cell(2, ci(column), field[3])
    for _, column, header in CYCLE_FIELDS:
        sheet.cell(2, ci(column), header)
    for r in (3, 4):
        sheet.cell(r, 1, f"synthetic-{r}")
        sheet.cell(r, ci("AD"), "SOX4 + 信迪利单抗")
        for col in GROUPED_COLUMNS:
            sheet.cell(r, ci(col), 1)
        for _, col, _ in CYCLE_FIELDS:
            sheet.cell(r, ci(col), 4)
        sheet.cell(r, ci("CK"), "forbidden-outcome")
    sheet["AD4"] = "NEVER-INTERPRET-TEST"
    path = tmp_path / "synthetic-grouped.xlsx"
    workbook.save(path)
    pseudonymizer = HMACPseudonymizer(bytes(range(32)))
    patient = pseudonymizer.token("patient", "synthetic-3", prefix="P")
    rows = read_regimen_rows(path, patient_ids={patient}, pseudonymizer=pseudonymizer)
    assert len(rows) == 1 and rows[0]["named_mentions"]["regimen_sox"] == 1
    assert "forbidden-outcome" not in str(rows) and "NEVER-INTERPRET" not in str(rows)
    sheet["AI3"] = 5
    sheet["AE3"] = 0
    workbook.save(path)
    rows = read_regimen_rows(path, patient_ids={patient}, pseudonymizer=pseudonymizer)
    assert rows[0]["cycles"]["chemotherapy"] is None
    assert rows[0]["cycle_quality"]["chemotherapy"] == "modality_count_conflict"
    sheet["AE2"] = "wrong code mapping"
    workbook.save(path)
    with pytest.raises(DataContractError) as error:
        read_regimen_rows(path, patient_ids={patient}, pseudonymizer=pseudonymizer)
    assert error.value.code == "REGIMEN_HEADER_MISMATCH"
