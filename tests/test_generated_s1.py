from __future__ import annotations

from dataclasses import asdict, replace

import pytest
import torch
from test_training import _batch as world_batch
from test_training import _model as world_model

from stageworld.data.baseline_clinical import (
    CLINICAL_FIELDS,
    FIELD_NAMES,
    encode_baseline,
    fit_clinical_transform,
    parse_baseline_fields,
    parse_clinical_value,
    read_baseline_rows,
)
from stageworld.data.paired_ct import HMACPseudonymizer
from stageworld.errors import DataContractError
from stageworld.generated_training import (
    GeneratedS1Trainer,
    GeneratedTrainingBatch,
    observed_supervision,
)
from stageworld.model.generated_s1 import READOUT_MODES, GeneratedS1Config, GeneratedS1Model
from stageworld.survival import piecewise_exponential_nll
from stageworld.training import LossWeights, TrainingPhase


def clinical_rows(count=4):
    return {
        f"person-{i}": parse_baseline_fields(
            {"age": 40 + 10 * i, "sex": "男" if i % 2 else "女", "ct_stage": "cT4a", "her2": i % 2}
        )
        for i in range(count)
    }


def generated_model(mode="history_generated", stochastic=True):
    values = asdict(world_model().config)
    values.update(model_version="stageworld-generated-s1-v1", use_stochastic_state=stochastic)
    return GeneratedS1Model(GeneratedS1Config(**values, readout_mode=mode)).eval()


def generated_batch(indices=(0, 1, 2, 3)):
    source = world_batch(indices)
    valid = source.survival_valid.clone()
    valid[:, 2] = False
    source = replace(
        source,
        survival_valid=valid,
        future_pathology_valid=torch.zeros_like(source.future_pathology_valid),
    )
    rows = clinical_rows(len(indices))
    transform = fit_clinical_transform(rows, set(rows))
    return GeneratedTrainingBatch.from_world(
        source, encode_baseline(list(rows.values()), transform)
    )


def test_clinical_transform_uses_training_patients_and_retains_missingness():
    rows = clinical_rows()
    transform = fit_clinical_transform(rows, {"person-0", "person-1"})
    assert transform["continuous"]["age"] == {"mean": 45.0, "scale": 5.0, "n": 2}
    data = encode_baseline(list(rows.values()), transform)
    age = FIELD_NAMES.index("age")
    assert data.values[:, age].tolist() == [-1.0, 1.0, 3.0, 5.0]
    assert not data.observed[:, FIELD_NAMES.index("eber")].any()
    assert torch.isfinite(data.ridge_features()).all()


@pytest.mark.parametrize(
    "name,value", [("lauren", 5), ("her2", 9), ("tps", 101), ("age", float("nan")), ("sex", 1)]
)
def test_invalid_or_unmeasured_clinical_values_are_not_negative_results(name, value):
    field = next(f for f in CLINICAL_FIELDS if f.name == name)
    assert parse_clinical_value(field, value)[0] is None


def test_baseline_query_rejects_future_or_identity_fields():
    for field in ("survival", "ypT", "ct1", "lauren_Z", "name"):
        with pytest.raises(DataContractError) as error:
            parse_baseline_fields({field: 1})
        assert error.value.code == "BASELINE_FIELD_NOT_ALLOWED"


def test_cn_positive_without_substage_stays_distinct_from_n1():
    field = next(f for f in CLINICAL_FIELDS if f.name == "cn_stage")
    assert parse_clinical_value(field, "N+") == ("+", "observed")
    assert parse_clinical_value(field, "N1") == ("1", "observed")


def test_workbook_reader_never_uses_ambiguous_lauren_or_later_columns(tmp_path):
    import openpyxl
    from openpyxl.utils import column_index_from_string as ci

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.cell(1, 5, "阶段1——患者基线数据")
    sheet.cell(2, 1, "序列号")
    for field in CLINICAL_FIELDS:
        sheet.cell(2, ci(field.column), field.header)
    sheet.cell(3, 1, "synthetic-one")
    sheet.cell(3, ci("H"), 60)
    sheet.cell(3, ci("Z"), "later-stage-answer")
    sheet.cell(3, ci("BV"), "later-stage-answer")
    path = tmp_path / "synthetic.xlsx"
    book.save(path)
    identity = HMACPseudonymizer(b"synthetic-only-test-key-no-real-data")
    patient = identity.token("patient", "synthetic-one", prefix="P")
    rows, audit = read_baseline_rows(path, {patient}, identity)
    assert rows[patient]["lauren"] is None
    assert rows[patient]["age"] == 60
    assert audit["ambiguous_lauren_Z_used"] is False


@pytest.mark.parametrize("mode", READOUT_MODES)
def test_prediction_needs_no_ct1_and_produces_valid_survival(mode, monkeypatch):
    model, batch = generated_model(mode), generated_batch()

    def forbidden(*args, **kwargs):
        raise AssertionError("Observation updater must not run during generated inference")

    monkeypatch.setattr(model, "update_posterior", forbidden)
    output = model(**batch.prediction_inputs(deterministic=True))
    prediction = output.survival_s1_pred
    assert prediction.simulated
    assert output.state_s1_pred.state_kind == "prior"
    assert torch.equal(prediction.query_time, batch.s1_time)
    assert prediction.risk.shape == (4, 3)
    assert torch.isfinite(prediction.risk).all()
    assert (prediction.risk[:, 1:] >= prediction.risk[:, :-1]).all()
    torch.testing.assert_close(prediction.risk + prediction.survival, torch.ones(4, 3))


def test_generated_prediction_is_invariant_to_ct1_and_outcomes():
    model, batch = generated_model(), generated_batch()
    changed = replace(
        batch,
        ct1=replace(batch.ct1, values=batch.ct1.values.flip(0) * 123),
        future_ct_target=batch.future_ct_target + 500,
        survival_durations=batch.survival_durations + 20,
    )
    original = model(**batch.prediction_inputs(deterministic=True))
    altered = model(**changed.prediction_inputs(deterministic=True))
    assert torch.equal(original.survival_s1_pred.rates, altered.survival_s1_pred.rates)
    assert torch.equal(original.state_s1_pred.memory, altered.state_s1_pred.memory)


def test_generated_survival_gradient_trains_transition_without_observation_update():
    model, batch = generated_model(), generated_batch()
    output = model(**batch.prediction_inputs(deterministic=True))
    loss = piecewise_exponential_nll(
        output.survival_s1_pred.rates.squeeze(-1),
        batch.survival_durations[:, 1],
        batch.survival_events[:, 1],
        model.survival_cutpoints,
    )
    loss.backward()
    assert any(p.grad is not None and p.grad.abs().sum() > 0 for p in model.transition.parameters())
    assert any(
        p.grad is not None and p.grad.abs().sum() > 0 for p in model.survival_decoder.parameters()
    )
    assert all(p.grad is None for p in model.updater.parameters())
    assert any(
        p.grad is not None and p.grad.abs().sum() > 0
        for p in model.baseline_clinical_encoder.parameters()
    )


def test_auxiliary_observation_does_not_move_generated_query_to_acquisition():
    model, batch = generated_model(), generated_batch()
    output = model(**batch.prediction_inputs(deterministic=True))
    auxiliary = observed_supervision(model, output, batch, deterministic=True)
    assert torch.equal(auxiliary.future_ct.target_time, batch.ct1_acquisition_time)
    assert torch.equal(auxiliary.pre_update.query_time, batch.ct1_availability_time)
    assert torch.equal(output.state_s1_pred.query_time, batch.s1_time)
    assert not torch.equal(output.future_ct.target_time, auxiliary.future_ct.target_time)


def test_history_only_readout_has_no_generated_memory_dependency():
    model, batch = generated_model("history_only"), generated_batch()
    output = model(**batch.prediction_inputs(deterministic=True))
    state = replace(output.state_s1_pred, memory=output.state_s1_pred.memory + 50)
    changed = model.read_survival(output.baseline, state, batch.treatment_actions, batch.horizons)
    assert torch.equal(changed.rates, output.survival_s1_pred.rates)


def test_patient_normalized_generated_training_uses_declared_loss_weights():
    model, batch = generated_model(), generated_batch()
    trainer = GeneratedS1Trainer(model, torch.optim.AdamW(model.parameters(), lr=2e-4))
    weights = LossWeights(survival=1, future_ct=0.1, future_pathology=0, kl=0.001)
    record = trainer.optimizer_step((batch,), phase=TrainingPhase.JOINT_SURVIVAL, weights=weights)
    c = record["components"]
    expected = (
        0.5 * (c["survival_s0"] + c["survival_s1"]) + 0.1 * c["future_ct"] + 0.001 * c["kl_ct"]
    )
    assert record["loss"] == pytest.approx(expected)
    assert record["effective_n"]["survival_s2"] == 0
    assert record["diagnostics"]["s1_reads_ct1"] == 0


@pytest.mark.parametrize("offset", [-1.0, float("nan")])
def test_invalid_target_time_is_rejected(offset):
    model, batch = generated_model(), generated_batch()
    inputs = batch.prediction_inputs(deterministic=True)
    inputs["target_time"] = torch.full_like(batch.s1_time, offset)
    with pytest.raises(DataContractError):
        model(**inputs)


def test_endpoint_summary_rollout_matches_chronological_replay_and_gradients():
    model, batch = generated_model(stochastic=False), generated_batch()
    context = model.initialize_baseline(
        batch.ct0, batch.baseline, batch.s0_time, deterministic=True
    )
    targets = batch.s1_time + torch.arange(batch.batch_size).float()
    times = targets[:, None].expand_as(batch.treatment_actions.event_time)
    actions = replace(batch.treatment_actions, event_time=times, available_time=times)
    fast = model.rollout_prior(context.state, actions, targets, deterministic=True)
    reference = model._replay_to(
        context.state, (actions,), targets, deterministic=True, generator=None
    )
    torch.testing.assert_close(fast.memory, reference.memory, rtol=1e-6, atol=1e-6)
    parameters = tuple(model.transition.parameters())
    fast_grad = torch.autograd.grad(fast.memory.square().sum(), parameters, retain_graph=True)
    ref_grad = torch.autograd.grad(reference.memory.square().sum(), parameters)
    for left, right in zip(fast_grad, ref_grad, strict=True):
        torch.testing.assert_close(left, right, rtol=1e-4, atol=1e-5)


def test_generated_padding_is_ignored_before_projection_and_has_finite_gradients():
    model, batch = generated_model(stochastic=False), generated_batch()
    ct_valid = batch.ct0.valid.clone()
    ct_valid[:, -1] = False
    clean_ct = replace(batch.ct0, valid=ct_valid)
    padded_ct = replace(
        clean_ct,
        values=clean_ct.values.masked_fill(~ct_valid[..., None], float("nan")),
        available_time=clean_ct.available_time.float().masked_fill(~ct_valid, float("nan")),
    )
    actions = batch.treatment_actions
    mask = actions.valid.clone()
    mask[0] = False
    clean_actions = replace(actions, valid=mask)
    dirty_actions = replace(
        clean_actions,
        values=actions.values.masked_fill(~mask[..., None], float("nan")),
        known_exposure=actions.known_exposure.masked_fill(~mask, float("nan")),
        event_time=actions.event_time.float().masked_fill(~mask, float("nan")),
    )
    inputs = batch.prediction_inputs(deterministic=True)
    inputs.update(ct0=clean_ct, scenario_actions=clean_actions)
    clean = model(**inputs)
    inputs.update(ct0=padded_ct, scenario_actions=dirty_actions)
    dirty = model(**inputs)
    torch.testing.assert_close(
        clean.survival_s1_pred.rates, dirty.survival_s1_pred.rates, rtol=0, atol=0
    )
    dirty.survival_s1_pred.rates.sum().backward()
    assert all(p.grad is None or torch.isfinite(p.grad).all() for p in model.parameters())


def test_generated_checkpoint_resume_keeps_clinical_and_readout_weights(tmp_path):
    from test_training import _metadata

    model, batch = generated_model(stochastic=False), generated_batch()
    trainer = GeneratedS1Trainer(model, torch.optim.AdamW(model.parameters(), lr=2e-4))
    weights = LossWeights(survival=0, future_ct=1, future_pathology=0, kl=0.01)
    trainer.optimizer_step([batch], phase=TrainingPhase.WORLD_PRETRAIN, weights=weights)
    metadata = replace(
        _metadata(),
        model_version=model.config.model_version,
        phase="world_pretrain",
        **{k: None for k in asdict(_metadata()) if k.startswith("parent_")},
    )
    trainer.save_checkpoint(tmp_path / "checkpoint.pt", metadata)
    trainer.optimizer_step([batch], phase=TrainingPhase.WORLD_PRETRAIN, weights=weights)
    resumed_model = generated_model(stochastic=False)
    resumed = GeneratedS1Trainer(
        resumed_model, torch.optim.AdamW(resumed_model.parameters(), lr=2e-4)
    )
    resumed.load_checkpoint(tmp_path / "checkpoint.pt", expected=metadata)
    resumed.optimizer_step([batch], phase=TrainingPhase.WORLD_PRETRAIN, weights=weights)
    for name, value in model.state_dict().items():
        assert torch.equal(value, resumed_model.state_dict()[name])


def test_generated_gradient_accumulation_matches_patient_batch():
    from stageworld.data.baseline_clinical import BaselineClinical

    model, batch = generated_model(stochastic=False), generated_batch()
    other = generated_model(stochastic=False)
    other.load_state_dict(model.state_dict())
    shards = []
    for start, indices in ((0, (0, 1)), (2, (2, 3))):
        shards.append(
            replace(
                generated_batch(indices),
                baseline=BaselineClinical(
                    batch.baseline.values[start : start + 2],
                    batch.baseline.categories[start : start + 2],
                    batch.baseline.observed[start : start + 2],
                ),
            )
        )
    weights = LossWeights(survival=1, future_ct=0.1, future_pathology=0, kl=0.001)
    for active, batches in ((model, [batch]), (other, shards)):
        trainer = GeneratedS1Trainer(active, torch.optim.SGD(active.parameters(), lr=0.01))
        trainer.optimizer_step(batches, phase=TrainingPhase.JOINT_SURVIVAL, weights=weights)
    for name, value in model.state_dict().items():
        torch.testing.assert_close(value, other.state_dict()[name], rtol=1e-5, atol=1e-6)
