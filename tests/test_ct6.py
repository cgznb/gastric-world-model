from __future__ import annotations

import json
from dataclasses import replace

import numpy as np
import pytest
import torch
from test_generated_s1 import generated_batch, generated_model
from typer.testing import CliRunner

from stageworld.artifacts import read_json
from stageworld.cache import CacheProvenance
from stageworld.cli import app
from stageworld.config import load_config
from stageworld.ct6_evaluation import seed_averaged_comparisons
from stageworld.ct6_training import EarlyStopState, make_trainer, train_ct6_phase, verify_phase
from stageworld.ct6_workflow import verify_portable, verify_smoke_recovery
from stageworld.data.baseline_clinical import (
    CT6_CLINICAL_SCHEMA,
    CT6_FIELD_NAMES,
    CT6_FIELDS,
    encode_baseline,
    fit_clinical_transform,
    parse_baseline_fields,
    read_baseline_rows,
)
from stageworld.data.paired_ct import HMACPseudonymizer
from stageworld.data.treatment_compact import (
    COMPACT_GROUPS,
    DESCRIPTOR_NAMES,
    compact_actions,
    encode_compact,
    fit_name_support,
    normalize_treatment,
    read_compact_rows,
    summarize_compact,
    treatment_fields,
)
from stageworld.data.treatment_regimens import GROUPED_COLUMNS, parse_named_mentions
from stageworld.data.treatment_summary import TREATMENT_FIELDS
from stageworld.errors import ArtifactError, DataContractError
from stageworld.generated_evaluation import nll_by_patient, predict_rates
from stageworld.generated_inference import (
    ct0_packet,
    export_inference_bundle,
    predict_generated_query,
)
from stageworld.generated_training import GeneratedS1Trainer
from stageworld.model.generated_s1 import GeneratedS1Model
from stageworld.real_workflow import RealFeatureBundle
from stageworld.survival import piecewise_exponential_nll
from stageworld.synthetic_workflow import _atomic_torch_save
from stageworld.training import TrainingPhase


def compact_row(patient="train", *, methods=None, regimens=None, drugs=None):
    return {
        "patient_id": patient,
        **normalize_treatment(
            {
                "methods": methods or {},
                "regimens": regimens or {},
                "drugs": drugs or {},
            }
        ),
    }


def ct6_rows(count=4):
    return {
        f"person-{i}": parse_baseline_fields(
            {"age": 40 + 10 * i, "bmi": 20 + i, "cn_stage": "N+"},
            schema_version=CT6_CLINICAL_SCHEMA,
        )
        for i in range(count)
    }


def ct6_batch(indices=(0, 1, 2, 3), *, unlabelled=False):
    batch = generated_batch(indices)
    rows = ct6_rows(len(indices))
    transform = fit_clinical_transform(rows, set(rows), schema_version=CT6_CLINICAL_SCHEMA)
    treatment = [compact_row(patient) for patient in batch.patient_ids]
    support = fit_name_support(treatment, set(batch.patient_ids))
    values, _ = encode_compact(treatment, support)
    return replace(
        batch,
        baseline=encode_baseline(list(rows.values()), transform),
        treatment_actions=compact_actions(values, batch.ct1_acquisition_time, provenance="test"),
        s1_time=batch.ct1_acquisition_time.clone(),
        ct1=replace(batch.ct1, available_time=batch.ct1.acquired_time.clone()),
        ct1_availability_time=batch.ct1_acquisition_time.clone(),
        survival_valid=torch.zeros_like(batch.survival_valid)
        if unlabelled
        else batch.survival_valid,
    )


def ct6_model(mode="history_generated"):
    return GeneratedS1Model(
        replace(
            generated_model(mode).config,
            action_input_dim=82,
            clinical_schema_version=CT6_CLINICAL_SCHEMA,
        )
    ).eval()


def test_ct6_normalization_missingness_and_training_only_fit():
    rows = ct6_rows()
    transform = fit_clinical_transform(
        rows, {"person-0", "person-1"}, schema_version=CT6_CLINICAL_SCHEMA
    )
    changed = {**rows, "person-3": {**rows["person-3"], "age": 90}}
    assert (
        fit_clinical_transform(
            changed, {"person-0", "person-1"}, schema_version=CT6_CLINICAL_SCHEMA
        )
        == transform
    )
    assert set(transform["continuous"]) == {"age", "bmi"}
    data = encode_baseline(list(rows.values()), transform)
    assert data.values.shape == (4, 6)
    assert data.values[:, 1].tolist() == [-1, 1, 3, 5]
    assert not data.observed[:, 0].any()
    assert rows["person-0"]["cn_stage"] == "+"
    assert torch.isfinite(data.ridge_features()).all()


@pytest.mark.parametrize(
    "field",
    [
        "lauren",
        "signet_ring",
        "differentiation",
        "her2",
        "mmr",
        "pdl1",
        "tps",
        "cps",
        "eber",
        "height",
        "weight",
        "location",
        "dentate_line",
        "survival",
        "ct1",
        "cycles",
    ],
)
def test_ct6_rejects_excluded_query_fields(field):
    with pytest.raises(DataContractError) as error:
        parse_baseline_fields({field: 1}, schema_version=CT6_CLINICAL_SCHEMA)
    assert error.value.code == "BASELINE_FIELD_NOT_ALLOWED"


def test_excel_excluded_fields_and_cycles_do_not_change_inputs(tmp_path):
    import openpyxl
    from openpyxl.utils import column_index_from_string as ci

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.cell(1, 5, "\u9636\u6bb51\u2014\u2014\u60a3\u8005\u57fa\u7ebf\u6570\u636e")
    sheet.cell(
        1,
        ci("AD"),
        "\u9636\u6bb52\u2014\u2014\u672f\u524d\u5316\u7597\u65b9\u6848\u3001\u5468\u671f\u53caCT\u8bc4\u4f30",
    )
    sheet.cell(2, 1, "\u5e8f\u5217\u53f7")
    sheet.cell(2, ci("AD"), "\u672f\u524d\u6cbb\u7597")
    for field in CT6_FIELDS:
        sheet.cell(2, ci(field.column), field.header)
    for col, (_, _, _, header) in zip(GROUPED_COLUMNS, TREATMENT_FIELDS, strict=True):
        sheet.cell(2, ci(col), header)
    sheet.cell(3, 1, "synthetic-only")
    sheet.cell(3, ci("H"), 60)
    sheet.cell(3, ci("X"), "N+")
    sheet.cell(3, ci("AD"), "SOX + FLOT + \u4fe1\u8fea\u5229\u5355\u6297")
    path = tmp_path / "synthetic.xlsx"
    book.save(path)
    identity = HMACPseudonymizer(b"synthetic-test-identity-key-no-private-rows")
    patient = identity.token("patient", "synthetic-only", prefix="P")
    before = read_baseline_rows(path, {patient}, identity, schema_version=CT6_CLINICAL_SCHEMA)
    treatment = read_compact_rows(path, {patient}, identity)
    for col in (
        "I",
        "J",
        "L",
        "M",
        "N",
        "O",
        "P",
        "Q",
        "R",
        "S",
        "T",
        "U",
        "V",
        "Z",
        "AH",
        "AI",
        "AJ",
        "BV",
    ):
        sheet.cell(2, ci(col), "unused-header")
        sheet.cell(3, ci(col), "=unusable_later_result()")
    sheet.cell(3, ci("AD"), "SOX + FLOT + \u4fe1\u8fea\u5229\u5355\u6297 99\u5468\u671f")
    book.save(path)
    assert before == read_baseline_rows(
        path, {patient}, identity, schema_version=CT6_CLINICAL_SCHEMA
    )
    assert treatment == read_compact_rows(path, {patient}, identity)
    assert tuple(before[0][patient]) == CT6_FIELD_NAMES
    assert treatment[0]["regimens"]["SOX"] == treatment[0]["regimens"]["FLOT"] == 1
    assert treatment[0]["drugs"]["sintilimab"] == 1
    assert "cycles" not in json.dumps(treatment)


def test_drug_identity_multiregimen_unknown_negative_and_group_layout():
    rows = [
        compact_row("a", regimens={"SOX": 1, "FLOT": 1}, drugs={"sintilimab": 1}),
        compact_row("b", drugs={"camrelizumab": 1}, methods={"hipec": 0}),
    ]
    support = fit_name_support(rows, {"a", "b"})
    values, flags = encode_compact(rows, support)
    assert values.shape == (2, 4, 82) and flags == ((), ())
    assert not torch.equal(values[0, 2], values[1, 2])
    assert values[0, 1, DESCRIPTOR_NAMES.index("regimen_sox")] == 1
    assert values[0, 1, DESCRIPTOR_NAMES.index("regimen_flot")] == 1
    hipec = DESCRIPTOR_NAMES.index("hipec")
    assert values[:, 0, hipec].tolist() == [0, 0]
    assert values[:, 0, 39 + hipec].tolist() == [0, 1]
    for g, indices in enumerate(COMPACT_GROUPS):
        outside = sorted(set(range(39)) - set(indices))
        assert not values[:, g, outside].any()
        assert not values[:, g, [39 + i for i in outside]].any()
        assert (values[:, g, 78:] == torch.eye(4)[g]).all()
    mentions = parse_named_mentions("SOX SOX", "s")
    assert mentions["regimen_sox"] == 1 and mentions["oxaliplatin"] is None


def test_zero_support_names_are_unknown_and_validation_cannot_fit_support():
    training = compact_row("train", drugs={"sintilimab": 1})
    validation = compact_row("validation", drugs={"camrelizumab": 1})
    support = fit_name_support([training, validation], {"train"})
    assert support == fit_name_support([training], {"train"})
    values, flags = encode_compact([validation], support)
    unknown, _ = encode_compact([compact_row()], support)
    assert torch.equal(values, unknown)
    assert flags == (("unseen_treatment_name:camrelizumab",),)
    assert summarize_compact([training, validation], support)["named_drug_count"] == 25


def test_conflicts_are_unknown_in_records_and_errors_in_queries():
    fields = {"methods": {"chemotherapy": 0}, "regimens": {"SOX": 1}, "drugs": {"oxaliplatin": 1}}
    with pytest.raises(DataContractError) as error:
        normalize_treatment(fields)
    assert error.value.code == "COMPACT_TREATMENT_CONFLICT"
    cleaned = normalize_treatment(fields, strict_conflicts=False)
    assert cleaned["conflicts"] == ["chemotherapy"]
    assert cleaned["methods"]["chemotherapy"] is None
    assert all(v is None for v in cleaned["regimens"].values())
    assert cleaned["drugs"]["oxaliplatin"] is None
    with pytest.raises(DataContractError) as error:
        normalize_treatment({"methods": {}, "regimens": {}, "drugs": {"invented": 1}})
    assert error.value.code == "COMPACT_TREATMENT_NAME"
    with pytest.raises(DataContractError) as error:
        normalize_treatment({**fields, "cycles": {}})
    assert error.value.code == "COMPACT_TREATMENT_FIELDS"


def test_ct6_future_invariance_and_survival_gradient():
    model, batch = ct6_model(), ct6_batch()
    output = model(**batch.prediction_inputs(deterministic=True))
    changed = replace(
        batch,
        ct1=replace(batch.ct1, values=batch.ct1.values.flip(0) * 999),
        future_ct_target=batch.future_ct_target + 123,
        survival_durations=batch.survival_durations + 10,
    )
    altered = model(**changed.prediction_inputs(deterministic=True))
    assert torch.equal(output.survival_s1_pred.rates, altered.survival_s1_pred.rates)
    loss = piecewise_exponential_nll(
        output.survival_s1_pred.rates.squeeze(-1),
        batch.survival_durations[:, 1],
        batch.survival_events[:, 1],
        model.survival_cutpoints,
    )
    loss.backward()
    assert any(p.grad is not None and p.grad.abs().sum() > 0 for p in model.transition.parameters())
    assert all(p.grad is None for p in model.updater.parameters())
    assert batch.future_ct_target.grad is None


def test_v2_portable_denies_other_files_and_rejects_mixed_schema(tmp_path):
    model, batch = ct6_model(), ct6_batch((0,))
    rows = ct6_rows(1)
    row = compact_row()
    support = fit_name_support([row], {"train"})
    export_inference_bundle(
        model,
        tmp_path / "inference.pt",
        clinical_snapshot={
            "artifact_id": "synthetic",
            "transform": fit_clinical_transform(
                rows, set(rows), schema_version=CT6_CLINICAL_SCHEMA
            ),
        },
        treatment_support=support,
        encoder_provenance=batch.ct0.provenance,
        weight_version="synthetic",
    )
    _atomic_torch_save(tmp_path / "ct0.pt", ct0_packet(batch.ct0))
    query = {
        "ct0_features": "ct0.pt",
        "baseline_clinical": rows["person-0"],
        "s0_time_days": 0,
        "target_interval_days": 30,
        "horizons_years": [1, 3],
        "treatment_scenario": {"description": "synthetic", **treatment_fields(row)},
    }
    (tmp_path / "query.json").write_text(json.dumps(query))
    assert verify_portable(tmp_path, predict_rates(model, [batch]))["status"] == "passed"
    output = read_json(tmp_path / "independent.json")
    assert output["cycle_counts_used"] is False and output["ct1_used"] is False
    cli = CliRunner().invoke(
        app,
        [
            "predict-generated-s1",
            "--config",
            str(tmp_path / "inference.json"),
            "--query-file",
            str(tmp_path / "query.json"),
            "--output-file",
            str(tmp_path / "cli.json"),
        ],
    )
    assert cli.exit_code == 0, cli.output
    assert read_json(tmp_path / "cli.json")["rates"] == output["rates"]
    config = read_json(tmp_path / "inference.json")
    config["schema_version"] = "generated-s1-portable-inference-v1"
    (tmp_path / "inference.json").write_text(json.dumps(config))
    with pytest.raises(DataContractError) as error:
        predict_generated_query(
            tmp_path / "inference.json", tmp_path / "query.json", tmp_path / "bad.json"
        )
    assert error.value.code == "GENERATED_CHECKPOINT_CONTRACT"


def test_early_stop_min_delta_and_ties_do_not_hide_best_weights():
    state = EarlyStopState()
    assert state.update(1.0, 1, 1e-4)
    assert state.update(0.99995, 2, 1e-4)
    assert state.selected_epoch == 2 and state.stale_epochs == 1
    assert not state.update(0.99995, 3, 1e-4)
    assert state.selected_epoch == 2 and state.stale_epochs == 2
    assert state.update(0.9998, 4, 1e-4) and state.stale_epochs == 0
    for epoch in range(5, 20):
        state.update(1.1, epoch, 1e-4)
    assert state.stale_epochs == 15
    with pytest.raises(ArtifactError):
        state.update(float("nan"), 20, 1e-4)


def tiny_phase_fixture():
    config = load_config("configs/project.generated-s1-ct6-drugs.yaml")
    config = replace(
        config,
        model=replace(
            config.model,
            hidden_dim=16,
            state_tokens=3,
            stochastic_dim_per_token=2,
            ct_tokens=2,
            ct_input_dim=8,
            pathology_input_dim=6,
            clinical_input_dim=4,
        ),
        training=replace(config.training, mixed_precision="off"),
        survival=replace(config.survival, finite_cutpoints=(0.0, 1.0, 3.0)),
    )
    batch = ct6_batch(unlabelled=True)
    validation = replace(batch, patient_ids=tuple(f"val-{i}" for i in range(4)))
    provenance = CacheProvenance.from_encoder(
        batch.ct0.provenance,
        schema_version="synthetic-v1",
        patch_sampling_version="synthetic-v1",
        split_version="synthetic-v1",
        target_transform_version="synthetic-v1",
        teacher_version="synthetic-v1",
    )
    bundle = RealFeatureBundle(
        {"train": (batch,), "validation": (validation,)},
        "synthetic-data",
        "synthetic-cohort",
        "synthetic-split",
        "synthetic-features",
        provenance,
        4,
        4,
        0,
    )
    return config, bundle, {"artifact_id": "synthetic-clinical"}


def assert_same_state(left, right):
    if isinstance(left, torch.Tensor):
        assert torch.equal(left, right)
    elif isinstance(left, dict):
        assert left.keys() == right.keys()
        for key in left:
            assert_same_state(left[key], right[key])
    elif isinstance(left, (list, tuple)):
        assert len(left) == len(right)
        for a, b in zip(left, right, strict=True):
            assert_same_state(a, b)
    else:
        assert left == right


@pytest.mark.parametrize("joint", [False, True])
def test_epoch_recovery_replays_model_optimizer_schedule_rng_and_counts(
    tmp_path, monkeypatch, joint
):
    config, bundle, snapshot = tiny_phase_fixture()
    args = dict(phase=TrainingPhase.WORLD_PRETRAIN, mode="history_generated", epochs=3)
    if joint:
        train_ct6_phase(config, bundle, snapshot, tmp_path / "parent", **args)
        parent = torch.load(tmp_path / "parent/selected.pt", weights_only=True)
        bundle = replace(
            bundle,
            batches_by_split={
                k: tuple(replace(b, survival_valid=ct6_batch().survival_valid) for b in batches)
                for k, batches in bundle.batches_by_split.items()
            },
        )
        args.update(phase=TrainingPhase.JOINT_SURVIVAL, parent=parent)
    uninterrupted = tmp_path / "uninterrupted"
    train_ct6_phase(config, bundle, snapshot, uninterrupted, **args)
    interrupted = tmp_path / "interrupted"
    original = GeneratedS1Trainer.optimizer_step

    def fail_in_second_epoch(self, *a, **kw):
        result = original(self, *a, **kw)
        if self.state.optimizer_step == 2:
            raise RuntimeError("synthetic interruption after uncommitted update")
        return result

    with monkeypatch.context() as patch:
        patch.setattr(GeneratedS1Trainer, "optimizer_step", fail_in_second_epoch)
        with pytest.raises(RuntimeError, match="synthetic interruption"):
            train_ct6_phase(config, bundle, snapshot, interrupted, **args)
    progress = read_json(interrupted / "progress.json")
    assert progress["completed_epochs"] == progress["optimizer_steps"] == 1
    assert progress["attempted_optimizer_steps"] == 2
    train_ct6_phase(config, bundle, snapshot, interrupted, resume=True, **args)
    expected = torch.load(uninterrupted / "final.pt", weights_only=True)
    actual = torch.load(interrupted / "final.pt", weights_only=True)
    for name in ("model_state", "optimizer_state", "scheduler_state", "rng_state", "trainer_state"):
        assert_same_state(expected[name], actual[name])
    assert read_json(uninterrupted / "history.json") == read_json(interrupted / "history.json")
    assert verify_phase(interrupted, args.get("parent"))["status"] == "passed"
    assert read_json(interrupted / "progress.json")["optimizer_steps"] == 3


def test_phase_plateau_stops_and_selected_final_are_distinct(tmp_path, monkeypatch):
    config, bundle, snapshot = tiny_phase_fixture()
    monkeypatch.setattr(
        "stageworld.ct6_training.feature_metrics", lambda *a: {"ct_loss": 1.0, "kl": 0.0}
    )
    _, report = train_ct6_phase(
        config,
        bundle,
        snapshot,
        tmp_path,
        phase=TrainingPhase.WORLD_PRETRAIN,
        mode="history_generated",
        epochs=100,
    )
    assert report["completed_epochs"] == report["optimizer_steps"] == 16
    assert report["stop_reason"] == "early_stopping" and report["selected_epoch"] == 1
    assert (
        report["verification"]["weight_versions"]["selected"]
        != report["verification"]["weight_versions"]["final"]
    )


def test_budget_interrupt_is_not_completed_training(tmp_path):
    config, bundle, snapshot = tiny_phase_fixture()
    config = replace(config, training=replace(config.training, development_max_minutes=0))
    with pytest.raises(ArtifactError) as error:
        train_ct6_phase(
            config,
            bundle,
            snapshot,
            tmp_path,
            phase=TrainingPhase.WORLD_PRETRAIN,
            mode="history_generated",
            epochs=100,
        )
    assert error.value.code == "CT6_TIME_BUDGET"
    assert read_json(tmp_path / "progress.json")["status"] == "interrupted"
    assert not (tmp_path / "final.pt").exists()


def test_smoke_recovery_probe_counts_discarded_update(tmp_path):
    config, bundle, snapshot = tiny_phase_fixture()
    reference = tmp_path / "reference"
    train_ct6_phase(
        config,
        bundle,
        snapshot,
        reference,
        phase=TrainingPhase.WORLD_PRETRAIN,
        mode="history_generated",
        epochs=2,
    )
    result = verify_smoke_recovery(config, bundle, snapshot, tmp_path / "probe", reference)
    assert result["status"] == "passed" and result["executed_optimizer_updates"] == 3


def test_optimizer_head_rates_decay_and_warmup(tmp_path):
    config, _, _ = tiny_phase_fixture()
    model = ct6_model()
    trainer = make_trainer(config, model, tmp_path, joint=True, steps_per_epoch=16, max_epochs=100)
    groups = {id(p): g for g in trainer.optimizer.param_groups for p in g["params"]}
    for name, parameter in model.named_parameters():
        head = name.startswith(("survival_decoder.", "survival_source."))
        assert groups[id(parameter)]["lr"] == pytest.approx(1e-5 if head else 2e-6)
        if "bias" in name or "embedding" in name or "norm" in name or "queries" in name:
            assert groups[id(parameter)]["weight_decay"] == 0
    assert trainer.scheduler is not None
    curve = trainer.scheduler.lr_lambdas[0]
    assert curve(0) == pytest.approx(0.1)
    assert curve(80) == pytest.approx(1)
    assert curve(1600) == pytest.approx(0.1)


def test_seed_averaged_bootstrap_uses_mean_of_metrics_and_support_rule():
    batch = ct6_batch()
    durations = torch.tensor([[0.7, 0.5, 0], [2.2, 2, 0], [3.2, 3, 0], [4.2, 4, 0]])
    events = torch.zeros(4, 3, dtype=torch.long)
    events[0, :2] = 1
    train = replace(batch, survival_durations=durations, survival_events=events)
    validation = replace(train, patient_ids=tuple(f"val-{i}" for i in range(4)))
    predictions = {
        s: {
            "history_generated": torch.full((4, 2, 2), rate),
            "history_only": torch.full((4, 2, 2), 0.2),
            "generated_only": torch.full((4, 2, 2), rate),
        }
        for s, rate in ((17, 0.1), (43, 0.4), (97, 0.8))
    }
    rows = seed_averaged_comparisons(predictions, [train], [validation], (0, 1, 3), replicates=20)
    nll = next(
        r
        for r in rows
        if r["stage"] == "S1_pred"
        and r["metric"] == "nll"
        and r["comparison"].endswith("history_only")
    )
    expected = np.mean(
        [
            float(
                (
                    nll_by_patient(p["history_generated"], [validation], (0, 1, 3))
                    - nll_by_patient(p["history_only"], [validation], (0, 1, 3))
                )[:, 1].mean()
            )
            for p in predictions.values()
        ]
    )
    assert nll["estimate_difference"] == pytest.approx(expected)
    for row in rows:
        if row["comparison"].endswith("generated_only") and row["estimate_difference"] is not None:
            assert row["estimate_difference"] == 0
    auc = next(
        r
        for r in rows
        if r["stage"] == "S1_pred" and r["metric"] == "cumulative_dynamic_auc" and r["horizon"] == 1
    )
    assert auc["confidence_lower"] is None
    assert auc["confidence_interval_status"] == "insufficient_valid_resamples"
